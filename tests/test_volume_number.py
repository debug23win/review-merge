"""Synthetic regression cases for old/new register columns and formatted reports."""
from pathlib import Path
from copy import copy
from datetime import datetime
import hashlib
import json
import subprocess
import uuid
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
MAIN = 'Все загруженные файлы'
TEMP = ROOT / 'tests/.runs' / uuid.uuid4().hex
TEMP.mkdir(parents=True)


def make(path, new=False, number=None):
    w = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx')
    s = w[MAIN]
    if new:
        s.insert_cols(13)
        s['M5'] = 'Номер\nтома'
        for col in range(13, 22):
            s.cell(6, col, col)
        if number is not None:
            s['M7'] = number
        s['V7'] = 'Пользовательский столбец'
        s.auto_filter.ref = 'A5:V11'
    w.save(path)
    return path


def run(target, *sources, success=True, legacy=False):
    completed = subprocess.run([str(ROOT / 'Свод_проверки.exe'), '--batch', str(target),
                                '2026-10-04', MAIN, *map(str, sources), *(['--keep-legacy-report'] if legacy else [])], timeout=90)
    assert (completed.returncode == 0) == success
    if success:
        return json.loads(Path(str(target) + '.run.json').read_text(encoding='utf-8-sig'))


def test_mixed_columns_and_text():
    target = make(TEMP / 'new-master.xlsx', True, '003.21.1.10')
    old = make(TEMP / 'old-source.xlsx')
    w = load_workbook(old)
    m = w[MAIN]
    for col in range(13, 19):
        m.cell(7, col, 1)
    m['S7'] = 'Замечание из старой описи'
    m['T7'] = '=Текст без вычисления'
    m['T7'].data_type = 's'
    m['C25'] = 'Новый документ'; m['D25'] = 'pdf'; m['H25'] = 'CAFEBABE'
    m['I25'] = 'Анна'; m['J25'] = datetime(2026, 10, 2); m['K25'] = 'ПД'; m['L25'] = 1
    m['R25'] = 1; m['T25'] = 'Новый текст'
    w.save(old)
    digest = hashlib.sha256(old.read_bytes()).hexdigest()
    report = run(target, old)
    w = load_workbook(target); m = w[MAIN]; ref = w['Справка по томам']
    assert m['M7'].value == '003.21.1.10' and m['M7'].number_format == '@'
    assert all(m.cell(7, c).value == 1 for c in range(14, 20))
    assert 'Замечание из старой описи' in m['T7'].value
    assert '=Текст без вычисления' in m['U7'].value and m['U7'].data_type != 'f'
    assert m['V7'].value == 'Пользовательский столбец' and m.auto_filter.ref.startswith('A5:W')
    assert m['W5'].value == 'Замечания устранены'
    appended = next(r for r in range(7, m.max_row + 1) if m.cell(r, 3).value == 'Новый документ')
    assert m.cell(appended, 19).value == 1 and m.cell(appended, 21).value == 'Новый текст'
    assert m.cell(appended, 13).number_format == '@'
    assert m['J7'].number_format.lower() == 'dd.mm.yyyy'
    h = next(c.column for c in ref[1] if c.value == 'ReviewMerge.FirstDates.v3')
    from openpyxl.utils import get_column_letter
    normalized = f'{get_column_letter(h+31)}7:{get_column_letter(h+36)}7'
    assert normalized in ref.cell(7, h + 5).value and '"?*"' not in ref.cell(7, h + 5).value
    assert f"'{MAIN}'!T7" in ref.cell(7, h + 23).value
    assert f"'{MAIN}'!U7" in ref.cell(7, h + 24).value
    assert all(f"'{MAIN}'!{chr(78 + n)}7" in ref.cell(7, h + 31 + n).value for n in range(6))
    assert hashlib.sha256(old.read_bytes()).hexdigest() == digest
    before = m['M7'].value, m['T7'].value, m['U7'].value
    run(target, old)
    m = load_workbook(target)[MAIN]
    assert before == (m['M7'].value, m['T7'].value, m['U7'].value)


def test_number_is_not_a_remark():
    source = make(TEMP / 'number-only.xlsx', True, '1')
    w = load_workbook(source); m = w[MAIN]
    for row in range(7, m.max_row + 1):
        for col in range(14, 22):
            m.cell(row, col).value = None
    m['M8'] = '01.02'; m['M9'] = '0001'; m['M10'] = '=A1'; m['M11'] = 'А/7-б'
    m['M10'].data_type = 's'
    w.save(source)
    target = TEMP / 'number-only-result.xlsx'
    assert run(target, source)['WithIssues'] == 0
    m = load_workbook(target)[MAIN]
    for row, value in [(7, '1'), (8, '01.02'), (9, '0001'), (10, '=A1'), (11, 'А/7-б')]:
        assert m.cell(row, 13).value == value and m.cell(row, 13).number_format == '@'
        assert m.cell(row, 13).data_type != 'f'


def test_number_conflict_and_legacy_target():
    target = make(TEMP / 'number-conflict.xlsx', True, '0007')
    source = make(TEMP / 'number-conflict-source.xlsx', True, '0008')
    report = run(target, source)
    assert report['Conflicts'] == 1 and load_workbook(target)[MAIN]['M7'].value == '0007'
    old = make(TEMP / 'legacy-master.xlsx')
    before = old.read_bytes()
    run(old, source)
    updated = load_workbook(old)[MAIN]
    assert updated['M5'].value == 'Номер тома' and updated['M7'].value == '0008'
    assert updated['M7'].number_format == '@'
    assert list(TEMP.glob('legacy-master.xlsx.backup-*'))[0].read_bytes() == before


def test_formatted_summary_is_preserved():
    target = make(TEMP / 'formatted.xlsx', True, '01.02')
    run(target, target, legacy=True)
    w = load_workbook(target); ref = w['Справка по томам']; h = next(c.column for c in ref[1] if c.value == 'ReviewMerge.FirstDates.v3')
    from openpyxl.utils import get_column_letter as letter
    file_last = max(c.row for row in ref for c in row if c.column == h + 11 and c.value)
    volume_last = max(c.row for row in ref for c in row if c.column == h + 12 and c.value)
    start = ref.cell(1, h + 3).value; end = ref.cell(1, h + 4).value
    w.remove(w['Свод']); s = w.create_sheet('Свод')
    s['A1'] = 'Оформленный пользователем свод'; s.merge_cells('A1:V1')
    s['A4'] = 'Этап'; s['B4'] = 'Дата / проверяющий'; s['A6'] = 'Этап 1'; s['B6'] = datetime(2026, 10, 4)
    s['H6'] = f'=COUNTIFS(\'Справка по томам\'!$F${start}:$F${end},">0")'
    s['J6'] = f'=COUNTIFS(\'Справка по томам\'!${letter(h+13)}$7:${letter(h+13)}${volume_last},">0",\'Справка по томам\'!${letter(h+2)}$7:${letter(h+2)}${volume_last},1)'
    s['C8'] = f'=COUNTA(\'Справка по томам\'!${letter(h+11)}$7:${letter(h+11)}${file_last})'
    s['C9'] = f'="\'Справка по томам\'!${letter(h+11)}$7:${letter(h+11)}${file_last}"'
    s['H6'].font = Font(name='Times New Roman', color='FF0000', bold=True)
    s['B6'].fill = PatternFill('solid', fgColor='AABBCC'); s.column_dimensions['B'].width = 38
    s.row_dimensions.group(7, 8, hidden=True); s.print_area = ['A1:V6', 'A8:V9']
    saved_style = copy(s['H6']._style); saved_merges = list(map(str, s.merged_cells.ranges)); saved_print = str(s.print_area)
    literal = s['C9'].value; w.save(target)
    source = make(TEMP / 'formatted-source.xlsx', False)
    w = load_workbook(source); m = w[MAIN]
    m['C30'] = 'Новый том'; m['D30'] = 'pdf'; m['H30'] = 'AABBCCDD'
    m['I30'] = 'Анна'; m['J30'] = datetime(2026, 10, 3); m['L30'] = 2
    w.save(source); run(target, source, legacy=True)
    w = load_workbook(target); s = w['СВОД']; ref = w['Справка по томам']
    assert s['A1'].value == 'Оформленный пользователем свод'
    assert s['H6']._style == saved_style and list(map(str, s.merged_cells.ranges)) == saved_merges
    assert str(s.print_area).upper() == saved_print.upper() and s.column_dimensions['B'].width == 38
    assert s.row_dimensions[7].hidden and s['B6'].fill.fgColor.rgb == '00AABBCC'
    assert s['C9'].value == literal, 'References inside quoted text must not change'
    h = next(c.column for c in ref[1] if c.value == 'ReviewMerge.FirstDates.v3')
    assert f'${letter(h+11)}${file_last + 1}' in s['C8'].value, s['C8'].value
    assert sum(bool(view.tabSelected) for sh in w for view in sh.views.sheetView) == 1


if __name__ == '__main__':
    test_mixed_columns_and_text()
    test_number_is_not_a_remark()
    test_number_conflict_and_legacy_target()
    test_formatted_summary_is_preserved()
    print('Volume column tests passed: mixed schemas, text identifiers, all six marks, both notes, append, conflicts, repeat merge, formulas and formatted summary')
