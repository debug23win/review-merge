"""Two-page PDF of the build, italic percentages of «СВОД» and brochures as parts of a volume; artificial data only."""
import json, re, subprocess, uuid, zlib
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MAIN = 'Все загруженные файлы'
B = ROOT / 'tests/.runs' / uuid.uuid4().hex
B.mkdir(parents=True)


def book(path, documents):
    w = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx'); m = w[MAIN]; m.delete_rows(7, m.max_row)
    for r, doc in enumerate(documents, 7):
        values = {1: r - 6, 4: 'pdf', 6: '01.10.2026', 7: '10:00', 8: f'{r:08X}', 11: 'ПД'}; values.update(doc)
        for c, v in values.items(): m.cell(r, c, v)
    w.save(path); return path


def pdf_pages(data):
    """Text of every page: glyph codes of the content streams decoded with the ToUnicode maps of their fonts."""
    objects = {int(m.group(1)): m.group(2) for m in re.finditer(rb'(\d+) 0 obj\n(.*?)\nendobj', data, re.S)}
    def stream(body):
        m = re.search(rb'stream\n(.*)\nendstream', body, re.S); raw = m.group(1)
        return zlib.decompress(raw) if b'/FlateDecode' in body[:m.start()] else raw
    maps = {}
    for number, body in objects.items():
        m = re.search(rb'/Subtype /Type0 .*?/ToUnicode (\d+) 0 R', body)
        if m:
            cmap = stream(objects[int(m.group(1))]).decode('ascii')
            maps[number] = {int(a, 16): chr(int(b, 16)) for a, b in re.findall(r'<([0-9A-F]{4})> <([0-9A-F]{4})>', cmap)}
    texts = []
    for number, body in sorted(objects.items()):
        if not re.search(rb'/Type /Page\b(?!s)', body): continue
        fonts = {name.decode(): maps[int(ref)] for name, ref in re.findall(rb'/(F\d+) (\d+) 0 R', body)}
        content = stream(objects[int(re.search(rb'/Contents (\d+) 0 R', body).group(1))]).decode('ascii')
        current, out = None, []
        for font, hexes in re.findall(r'/(F\d+) [\d.]+ Tf[^<]*<([0-9A-F]*)> Tj', content):
            out.append(''.join(fonts[font].get(int(hexes[i:i + 4], 16), '?') for i in range(0, len(hexes), 4)))
        texts.append((re.search(rb'/MediaBox \[([^\]]*)\]', body).group(1).decode(), ' '.join(out)))
    return texts, data


documents = [{3: 'Том 1', 9: 'Анна', 10: datetime(2026, 10, 2), 12: 1},
             {3: 'Том 2 (Брошюра 1 из 2)', 9: 'Анна', 10: datetime(2026, 10, 2), 12: 1},
             {3: 'Том 2 (Брошюра 2 из 2)', 9: 'Анна', 10: datetime(2026, 10, 2), 12: 1},
             {3: 'Том 3', 9: 'Борис', 10: datetime(2026, 10, 3), 12: 2}, {3: 'Том 4', 12: 2}]
output = book(B / 'svod.xlsx', documents)
subprocess.run([str(ROOT / 'Свод_проверки.exe'), '--batch', str(output), '2026-10-04', MAIN, str(output)], check=True, timeout=120)
report = json.loads(Path(str(output) + '.run.json').read_text(encoding='utf-8-sig'))
pdf = Path(report['Pdf'])
assert report['PdfError'] is None and pdf.name == 'svod 04.10.2026.pdf' and pdf.parent == output.parent, report

cached = load_workbook(output, data_only=True)
work = cached['Итоги работы']
assert work.cell(6, 2).value == 'Всего' and work.cell(6, work.max_column).value == 3, 'Two brochures are one volume: Том 1, Том 2, Том 3'

summary = load_workbook(output)['СВОД']
for col in range(5, 20, 2):
    font = summary.cell(7, col).font
    assert font.i and font.sz == 9, (col, font.i, font.sz, 'percentages are italic, size 9')
    count = summary.cell(7, col - 1).font
    assert not count.i and count.sz == 12, (col - 1, count.i, count.sz, 'counts keep their font')
assert summary.cell(7, 9).font.color.rgb == summary.cell(7, 8).font.color.rgb, 'percentages keep the color of their block'

pages, data = pdf_pages(pdf.read_bytes())
assert data.startswith(b'%PDF-1.7') and data.rstrip().endswith(b'%%EOF')
assert len(pages) == 2, len(pages)
assert all(box == '0 0 842 595' for box, _ in pages), [box for box, _ in pages]
assert b'/FontFile2' in data and b'/ToUnicode' in data
first, second = pages[0][1], pages[1][1]
assert 'СВОДНАЯ ТАБЛИЦА ПО РЕЗУЛЬТАТАМ ПРОВЕРКИ ДОКУМЕНТАЦИИ' in first and 'Проверено за сутки' in first and '04.10.2026' in first, first[:300]
assert 'ИТОГАМ РАБОТЫ' not in first and 'Альтернативное' not in first
assert 'СВОДНАЯ ТАБЛИЦА ПО ИТОГАМ РАБОТЫ ПРОВЕРЯЮЩИХ' in second and 'Анна' in second and 'Борис' in second and 'ВСЕГО' in second, second[:300]
assert 'Альтернативное' not in second, 'Only the first table of «Итоги работы» is on the second page'

# The PDF alone can be made again from the saved workbook.
pdf.unlink()
subprocess.run([str(ROOT / 'Свод_проверки.exe'), '--pdf', str(output), '2026-10-04'], check=True, timeout=120)
assert len(pdf_pages(pdf.read_bytes())[0]) == 2
print('PDF report tests passed: two pages, summary and work totals, embedded fonts, italic percentages and brochures')
