"""Independent PD regression cases: resolved marks, user columns, dates and grouping."""
import json, re, subprocess, uuid, shutil, os, zipfile
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'tests/.runs' / uuid.uuid4().hex
BASE.mkdir(parents=True)
MAIN = 'Все загруженные файлы'

def fixture(name, records, resolved=False):
    path = BASE / name
    w = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx')
    m = w[MAIN]; m.delete_rows(7, m.max_row)
    if resolved: m['U5'] = 'Замечания устранены'; m['U6'] = 21
    for r, data in enumerate(records, 7):
        values = {1:r-6, 3:'Том '+str(r-6), 4:'pdf', 6:'01.10.2026', 7:'10:00', 8:f'{r:08X}', 9:'Анна', 10:datetime(2026,10,2), 11:'ПД', 12:1}
        values.update(data)
        for c, value in values.items(): m.cell(r,c,value)
    w.save(path)
    return path

def build(path, *sources):
    subprocess.run([str(ROOT/'Свод_проверки.exe'),'--batch',str(path),'2026-10-04',MAIN,str(path),*map(str,sources)],check=True,timeout=180)
    return json.loads(Path(str(path)+'.run.json').read_text(encoding='utf-8-sig'))

def summary(path):
    s=load_workbook(path,data_only=True)['СВОД']
    return tuple(int(re.search(r': (\d+)$',s.cell(4,c).value).group(1)) for c in (4,8,12))

def resolved_column(m):
    return next(c.column for c in m[5] if c.value=='Замечания устранены')

def test_marks():
    marks=['0.0','-0','0,00','0e3',False,'\u00a0\t','нет',1,datetime(2026,10,4),'да',True]
    path=fixture('marks.xlsx',[{13:1,21:mark} for mark in marks],True)
    build(path)
    assert summary(path)==(11,4,7), ('resolved zeros and booleans',summary(path))
    build(path)
    assert summary(path)==(11,4,7), 'repeat merge must preserve the same resolved count'
    return path

def test_user_columns():
    path=fixture('user-columns.xlsx',[{13:1},{13:1}])
    w=load_workbook(path);m=w[MAIN]
    m['U7']='Сохранить пользовательский текст'
    m['V8']='=1+1'
    w.save(path);build(path)
    w=load_workbook(path);m=w[MAIN]
    assert m['V7'].value=='Сохранить пользовательский текст' and m['W8'].value=='=1+1'
    assert resolved_column(m)==24, 'new header must follow data and formulas, including columns without a header'
    assert summary(path)==(2,0,2)
    build(path)
    w=load_workbook(path);m=w[MAIN]
    assert m['V7'].value=='Сохранить пользовательский текст' and m['W8'].value=='=1+1'
    assert resolved_column(m)==24 and summary(path)==(2,0,2)

def test_changed_date_format():
    path=fixture('mark-format.xlsx',[{13:1,21:datetime(2026,10,3)}],True)
    build(path);w=load_workbook(path);m=w[MAIN];c=m.cell(7,resolved_column(m));c.value=1;w.save(path)
    build(path);m=load_workbook(path)[MAIN];c=m.cell(7,resolved_column(m))
    assert c.value==1 and not c.is_date, (c.value,c.number_format)

def test_invalid_marks_and_whitespace():
    path=fixture('invalid-marks.xlsx',[{13:'+',21:1},{13:2,21:1},{19:'\u00a0\t ',21:1},{13:1,21:1}],True)
    report=build(path)
    assert report['WithIssues']==1 and summary(path)==(1,1,0), (report['WithIssues'],summary(path))

def test_rules():
    compiler=Path(os.environ['WINDIR'])/'Microsoft.NET/Framework64/v4.0.30319/csc.exe'
    if not compiler.exists(): compiler=Path(os.environ['WINDIR'])/'Microsoft.NET/Framework/v4.0.30319/csc.exe'
    target=BASE/'PdRulesTest.exe'
    subprocess.run([str(compiler),'/nologo','/codepage:65001','/out:'+str(target),'/r:'+str(ROOT/'Свод_проверки.exe'),str(ROOT/'tests/PdRulesTest.cs')],check=True)
    shutil.copy2(ROOT/'Свод_проверки.exe',BASE/'Свод_проверки.exe')
    subprocess.run([str(target)],check=True,timeout=20)

def test_fragments():
    path=fixture('fragments.xlsx',[{3:'Том А фрагмент 1',13:1},{3:'Том А фрагмент 2',13:1,21:1},{3:'Том А-УЛ'},
                                   {3:'Том Б',13:1},{3:'Том В',21:1}],True)
    build(path)
    assert summary(path)==(2,1,1),summary(path)
    return path

def helper(path):
    """«Справка по томам» (formulas and cached values) and the first helper column."""
    w,c=load_workbook(path),load_workbook(path,data_only=True)
    s,v=w['Справка по томам'],c['Справка по томам']
    h=next(x.column for x in s[1] if x.value=='ReviewMerge.FirstDates.v3')
    return s,v,h

def formula_limits(path):
    """Longest formula and largest argument count of one function call in the book."""
    longest=arguments=0
    for s in load_workbook(path).worksheets:
        for row in s.iter_rows():
            for c in row:
                if not(isinstance(c.value,str) and c.value.startswith('=')): continue
                f=re.sub(r'"(?:[^"]|"")*"','""',c.value);longest=max(longest,len(c.value));stack=[]
                for ch in f:
                    if ch=='(': stack.append(1)
                    elif ch==')': arguments=max(arguments,stack.pop())
                    elif ch==',' and stack: stack[-1]+=1
    return longest,arguments

def test_raw_file_name_and_many_parts():
    # The file-name helper caches the cell as the formula returns it; a volume of 300 parts keeps Excel's limits.
    records=[{3:'Том 6 '},{3:12345}]
    for k in range(1,301): records+=[{3:'Том-А Фрагмент %d'%k,13:1 if k==7 else None},{3:'Том-Б%d'%k,12:2}]
    path=fixture('many-parts.xlsx',records)
    report=build(path)
    s,v,h=helper(path)
    names={v.cell(r,h+11).value for r in range(7,v.max_row+1)}
    assert 'Том 6 ' in names and 12345 in names, 'raw file names'
    longest,arguments=formula_limits(path)
    assert longest<=8192 and arguments<=255, (longest,arguments)
    row=next(r for r in range(7,s.max_row+1) if 'AGGREGATE' in str(s.cell(r,h+13).value))
    assert v.cell(row,h+13).value==datetime(2026,10,2) and v.cell(row,h+14).value==1, 'first check and remarks of the volume'
    assert 'SUMPRODUCT' in s.cell(row,h+13).value
    assert report['Reviewed']==602 and summary(path)[0]==1

def test_box_text_and_criteria():
    # Line breaks and no-break spaces are spaces; letters only give no box; «*» and long box lists are not COUNTIF patterns.
    boxes=['45\n','45 и 46\n',' 47, 48','б/н','100-145']
    path=fixture('boxes.xlsx',[{12:b,9:'*' if i==2 else 'Анна'} for i,b in enumerate(boxes)])
    build(path)
    s,v,h=helper(path)
    assert [v.cell(r,h).value or '' for r in range(7,12)]==[';ПД|45;',';ПД|45;ПД|46;',';ПД|47;ПД|48;','',';'+';'.join('ПД|%d'%n for n in range(100,146))+';']
    assert [v.cell(r,h+30).value for r in range(7,10)]==['45','45,46','47,48']
    assert 'CHAR(10)' in s.cell(7,h+30).value and '"б/н",""' in s.cell(10,h).value.replace('\n',' ')
    criteria=[s.cell(r,c).value for r in range(12,16) for c in range(2,12) if isinstance(s.cell(r,c).value,str)]
    assert criteria and all('"~*"' in f for f in criteria if 'COUNTIF' in f)
    row,col=next((r,c) for r in (12,14) for c in range(2,20) if len(str(v.cell(r,c).value or ''))>255)
    assert 'LEN(' in s.cell(row+1,col).value and 'SUMPRODUCT' in s.cell(row+1,col).value and v.cell(row+1,col).value==1

def test_reference_sheet_xml():
    # A user-made «Справка по томам» without <dimension> and with an autoFilter builds in schema order.
    path=fixture('xml.xlsx',[{},{12:2}])
    w=load_workbook(path);s=w.create_sheet('Справка по томам');s['A1'],s['B1']='Коробов получено',10;s['A12']='Анна';s.auto_filter.ref='A51:H60';w.save(path)
    with zipfile.ZipFile(path) as z: items={n:z.read(n) for n in z.namelist()}
    name=next(n for n in items if n.startswith('xl/worksheets/') and b'A51:H60' in items[n])
    items[name]=re.sub(rb'<dimension [^>]*/>',b'',items[name])
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in items.items(): z.writestr(n,b)
    build(path)
    with zipfile.ZipFile(path) as z: xml=z.read(name).decode('utf-8')
    order=re.findall(r'<(dimension|sheetData|autoFilter|mergeCells|pageMargins)[ >/]',xml)
    assert order==['dimension','sheetData','autoFilter','mergeCells','pageMargins'], order

def test_standalone_iul_remark_types():
    # A standalone ИУЛ is not a volume, but its remark types are counted.
    path=fixture('iul.xlsx',[{3:'Шифр-СМ2-УЛ-ГТСС1',17:1},{3:'Том 2',12:2}])
    build(path)
    w,c=load_workbook(path)['Стат. по типам'],load_workbook(path,data_only=True)['Стат. по типам']
    column=next(x.column for x in w[4] if x.value=='Несоответствие контрольной суммы в ИУЛ')
    row=next(r for r in range(6,w.max_row+1) if c.cell(r,2).value==datetime(2026,10,2))
    assert c.cell(row,column).value==1 and c.cell(row,19).value==1, (c.cell(row,column).value,c.cell(row,19).value)

def test_box_table_layout_and_user_cells():
    # «Всего томов по описи» is found by its header; user text beside and below the box table survives a rebuild and its growth.
    path=fixture('table.xlsx',[{12:1},{12:2},{12:3}])
    build(path)
    s,v,h=helper(path)
    header=next(r for r in range(12,s.max_row+1) if s.cell(r,1).value=='Вид' and s.cell(r,2).value=='Короб')
    w=load_workbook(path);s=w['Справка по томам']
    for i in range(3): s.cell(header+1+i,5).value=10+i
    s.cell(header+1,10).value='рядом с коробом';end=int(v.cell(1,h+4).value);s.cell(end+2,1).value='Примечание под таблицей'
    s.delete_cols(3);w.save(path)
    source=fixture('more.xlsx',[{12:1},{12:2},{12:3}]+[{3:'Доп %d'%i,12:10+i} for i in range(8)])
    build(path,source)
    s,v,h=helper(path)
    header=next(r for r in range(12,s.max_row+1) if s.cell(r,1).value=='Вид' and s.cell(r,2).value=='Короб')
    assert s.cell(header,5).value=='Всего томов по описи' and [v.cell(header+1+i,5).value for i in range(3)]==[10,11,12]
    texts={c.value:(c.row,c.column) for row in s.iter_rows(max_col=h-1) for c in row if c.value in ('рядом с коробом','Примечание под таблицей')}
    end=int(v.cell(1,h+4).value)
    assert texts.get('рядом с коробом',(0,0))[1]==9 and texts.get('Примечание под таблицей',(0,0))[0]>end, texts
    # A box number changed later: the register count is live, the stale box row does not count.
    assert '$3' in load_workbook(path)['СВОД']['D8'].value and v.cell(3,h+44).value==11

def test_old_summary_calendar_and_long_names():
    # Dates of an old «Свод» after the report date are not days of the recreated reports; long names get taller rows.
    path=fixture('old.xlsx',[{9:'Яковлева Александра Алексеевна'},{9:'Яковлева Богдана Викторовна',12:2}])
    w=load_workbook(path);s=w.create_sheet('Свод');s['A7'],s['A8']='Получено коробов','Кол-во томов по акту'
    for i,d in enumerate((datetime(2026,10,1),datetime(2026,10,2),datetime(2026,10,9))): s.cell(5,2+i,d)
    w.save(path);build(path)
    days=lambda: [c.value for c in load_workbook(path)['СВОД']['B'] if isinstance(c.value,datetime)]
    first=days();assert datetime(2026,10,9) not in first, first
    build(path);assert days()==first
    w=load_workbook(path)
    assert w['Итоги работы'].row_dimensions[4].height>24 and w['СВОД'].row_dimensions[9].height>24

if __name__=='__main__':
    path=test_marks()
    test_user_columns()
    test_changed_date_format()
    test_invalid_marks_and_whitespace()
    test_rules()
    test_fragments()
    test_raw_file_name_and_many_parts()
    test_box_text_and_criteria()
    test_reference_sheet_xml()
    test_standalone_iul_remark_types()
    test_box_table_layout_and_user_cells()
    test_old_summary_calendar_and_long_names()
    print('PD recheck passed: resolved zero variants, boolean marks, preserved user columns, repeat merge, numeric format, invalid marks, '
          'raw file names, 300-part volumes, box text, escaped criteria, reference sheet XML, standalone IUL types, box table and old summary days')
    print(path)
