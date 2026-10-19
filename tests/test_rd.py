"""Working documentation (--rd): the RD check sheet «Проверка комплектов» merged from reviewers' copies; artificial data only."""
import json, shutil, subprocess, uuid
from datetime import datetime, date, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXE = ROOT / 'Свод_проверки.exe'
MAIN = 'Проверка комплектов'
B = ROOT / 'tests/.runs' / uuid.uuid4().hex
B.mkdir(parents=True)
FLAGS = ['П1', 'П2', 'П3', 'П4', 'П5', 'П6', 'Э1', 'Э2', 'Э3', 'Э4', 'Э5', 'Э6', 'ПР']
HEADS = ['Номер элемента', 'Наименование элемента', 'Шифр ПД', 'Шифр РД', 'Наименование комплекта', 'Изм.', 'Раздел', 'Сверка с ВЭ РД', 'Согласование ТЭ', 'Дата ТЭ', '№ письма ТЭ',
         'Проверил', 'Дата проверки'] + [f + '. вид замечания' for f in FLAGS] + ['Замечания (текст для письма подрядчику)', 'Примечание', 'Факт СМР', 'Видов замечаний', 'Замечания устранены',
         'Статус комплекта', 'Действующая редакция', 'Порядковый № элемента', 'Элемент строки']
columns = lambda heads: {h.split('.')[0] if h[:2] in FLAGS else h: i for i, h in enumerate(heads, 1)}
COL = columns(HEADS)
# The form with the deadline typed by the operator and the priority after «Дата проверки», as columns N and O of the stage 7 book.
HEADS_DUE = HEADS[:13] + ['Дата завершения проверки', 'Приоритет проверки '] + HEADS[13:]
# The stage list: element, name, sets with the status of ТЭ.
STAGE = [('1.1', 'Мост', [('К1', 'Да'), ('К2', 'Да'), ('К3', 'Замечания')]), ('1.2', 'Труба', [('Т1', 'Да'), ('Т2', 'Да')]), ('1.3', 'Дорога', [('Д1', 'Да')])]
NAME = {el: name for el, name, _ in STAGE}
LETTER = {'К3': datetime(2026, 10, 6)}  # the date of the letter of ТЭ («Дата», column K of «Ведомость»)
D = lambda day: datetime(2026, 10, day)


def form(path, rows, stage=True, acts=(), formula_rows=0, extra=None, heads=HEADS):
    """An RD check book: the check sheet (header in row 5, column numbers in row 6, data from row 7), «Ведомость» and «Проверка актов»."""
    wb = Workbook(); m = wb.active; m.title = MAIN; col = columns(heads)
    m['A1'] = 'Форма проверки рабочей документации'; m['A2'] = '9 этап – «Испытательный участок»'
    for c, h in enumerate(heads, 1): m.cell(5, c, h); m.cell(6, c, c)
    for r, row in enumerate(rows, 7):
        for key, value in row.items(): m.cell(r, col[key], value)
    for r in range(7, 7 + formula_rows): m.cell(r, col['Статус комплекта'], '=""')  # the formula rows a form has prepared
    if stage:
        v = wb.create_sheet('Ведомость')
        for c, h in {1: 'Номер элемента', 2: 'Наименование элемента', 3: 'Шифр ПД', 4: 'Шифр РД', 5: 'Наименование документации', 9: 'Статус комплекта'}.items(): v.cell(5, c, h)
        for c in range(1, 31): v.cell(6, c, c)
        r = 7
        for el, name, sets in STAGE:
            for k, (code, te) in enumerate(sets):
                if k == 0: v.cell(r, 1, el); v.cell(r, 2, name)
                v.cell(r, 4, code); v.cell(r, 5, 'Комплект ' + code); v.cell(r, 9, te); v.cell(r, 11, LETTER.get(code)); r += 1
    a = wb.create_sheet('Проверка актов')
    for c, h in enumerate(['№ п/п', 'Номер элемента', 'Наименование элемента', '№ акта', 'Дата акта', 'Проверил', 'Дата проверки', 'А1. акт', 'А2. акт', 'Замечания (текст)', 'Замечания устранены', 'Результат проверки акта'], 1):
        a.cell(5, c, h); a.cell(6, c, c)
    for r, act in enumerate(acts, 7):
        for c, value in act.items(): a.cell(r, c, value)
    if extra: extra(wb)
    wb.save(path); return path


def row(code, el=None, **values):
    data = {'Шифр РД': code}
    if el: data.update({'Номер элемента': el, 'Наименование элемента': NAME[el]})
    data.update({{'who': 'Проверил', 'date': 'Дата проверки', 'text': 'Замечания (текст для письма подрядчику)', 'note': 'Примечание', 'fixed': 'Замечания устранены', 'rev': 'Изм.', 'due': 'Дата завершения проверки'}.get(k, k): v for k, v in values.items()})
    return data


def build(output, *sources, day='2026-10-08', switches=()):
    result = subprocess.run([str(EXE), '--batch', str(output), day, MAIN, *map(str, sources), '--rd', *switches], timeout=300)
    error = Path(str(output) + '.error.txt')
    if error.exists(): return {'error': error.read_text(encoding='utf-8-sig')}
    assert result.returncode == 0
    return json.loads(Path(str(output) + '.run.json').read_text(encoding='utf-8-sig'))


def assert_sums(path):
    """The reports add up: days to totals, reviewers to the date rows, «Итоги проверки» and «Статистика» to «СВОД»."""
    wb = load_workbook(path, data_only=True); sv, tv, st = wb['СВОД'], wb['Итоги проверки'], wb['Статистика']
    num = lambda v: v if isinstance(v, (int, float)) else 0
    rows = [r for r in range(8, sv.max_row + 1) if isinstance(sv.cell(r, 2).value, datetime)]
    people = rows[1] - rows[0] - 1 if len(rows) > 1 else sv.max_row
    report = rows[-1]
    for k, r in enumerate(rows):
        persons = range(r + 1, r + 1 + people) if len(rows) > 1 else []
        for c in (8, 10, 16, 18):
            if persons: assert num(sv.cell(r, c).value) == sum(num(sv.cell(p, c).value) for p in persons), ('reviewers to the date row', r, c)
        for total, daily in ((8, 16), (10, 18)):
            before = num(sv.cell(rows[k - 1], total).value) if k else 0
            assert num(sv.cell(r, total).value) == before + num(sv.cell(r, daily).value), ('cumulative and daily', r, total)
            assert num(sv.cell(r, total - 4).value) - num(sv.cell(r, total).value) == num(sv.cell(r, total + 4).value), ('remaining', r, total)
    width = next(c for c in range(3, tv.max_column + 1) if tv.cell(4, c).value == 'ВСЕГО') + 1
    last = max(r for r in range(7, tv.max_row + 1) if isinstance(tv.cell(r, 2).value, datetime) and r < 7 + len(rows) + 1)
    for c in range(3, width + 1):
        assert num(tv.cell(6, c).value) == sum(num(tv.cell(r, c).value) for r in range(7, last + 1)), ('days to «Всего»', c)
    for r in range(6, last + 1):
        for j in (0, 1): assert num(tv.cell(r, width - 1 + j).value) == sum(num(tv.cell(r, c).value) for c in range(3 + j, width - 1, 2)), ('reviewers to «ВСЕГО»', r)
    assert (num(tv.cell(6, width - 1).value), num(tv.cell(6, width).value)) == (num(sv.cell(report, 8).value), num(sv.cell(report, 10).value)), '«Итоги проверки» and «СВОД»'
    stats = {st.cell(r, 1).value: st.cell(r, 2).value for r in range(6, st.max_row + 1) if st.cell(r, 1).value}
    assert stats['Проверено по состоянию на дату отчёта (как в «СВОД»)'] == num(sv.cell(report, 10).value), '«Статистика» and «СВОД»'
    statuses = ['Готово к ВПР', 'С неустранёнными замечаниями', 'Не согласовано ТЭ', 'Не соответствуют ВЭ РД (нет в ведомости или другой элемент)']
    assert sum(num(stats[k]) for k in statuses) <= num(sv.cell(report, 6).value)


target = form(B / 'Сводная РД.xlsx', [row('К1', '1.1', who='Анна', date=D(5)), row('К2', who='Анна', date=D(5), П1=1, text='Нет подписи'),
                                       row('К3', who='тех. эксперт', ПР=1, text='Не согласован ТЭ')],
              acts=[{2: '1.2', 4: 'А-1', 6: 'Борис', 7: D(8), 8: 1}])
boris = form(B / 'Борис.xlsx', [row('Т1', '1.2', who='Борис', date=D(6)), row('Т2', who='Борис', date=D(6), Э3=1), row('К2', '1.1', who='Анна', date=D(5), fixed=D(7), text='Нет подписи')])
vera = form(B / 'Вера.xlsx', [row('Д1', '1.3', who='Вера', date=D(7), note='проверить ВОР'), row('Т2', '1.2', who='Вера', date=D(7), П2=1, text='Нет печати')])
report = build(target, boris, vera)
assert 'error' not in report, report
assert (report['Rd'], report['Documents'], report['Elements'], report['Boxes'], report['Reviewed']) == (True, 6, 3, 3, 6), report
assert report['Pdfs'] and all(Path(p).exists() for p in report['Pdfs']) and report['PdfError'] is None, report
assert [Path(p).name for p in report['Pdfs']] == ['Сводная РД СВОД 2026.10.08.pdf', 'Сводная РД Итоги проверки 2026.10.08.pdf'], report['Pdfs']

wb, cached = load_workbook(target), load_workbook(target, data_only=True)
assert {'СВОД', 'Итоги проверки', 'Статистика', 'Элементы', 'Ведомость', 'Проверка актов', MAIN} <= set(wb.sheetnames), wb.sheetnames
assert wb.active.title == 'СВОД'
m, mv = wb[MAIN], cached[MAIN]
codes = {m.cell(r, COL['Шифр РД']).value: r for r in range(7, m.max_row + 1) if m.cell(r, COL['Шифр РД']).value}
assert list(codes) == ['К1', 'К2', 'К3', 'Т1', 'Т2', 'Д1'], codes
get = lambda code, key: m.cell(codes[code], COL[key]).value
# New sets go after the data in file order; an element number is written only where the element changes.
assert [get(c, 'Номер элемента') for c in ('Т1', 'Т2', 'Д1')] == ['1.2', None, '1.3'] and get('Т1', 'Наименование элемента') == 'Труба'
# Every «1» is kept, the reviewer and the date come from the latest dated check, «Замечания устранены» from the files.
assert (get('Т2', 'Проверил'), get('Т2', 'Дата проверки'), get('Т2', 'П2'), get('Т2', 'Э3'), get('Т2', 'Замечания (текст для письма подрядчику)')) == ('Вера', D(7), 1, 1, 'Нет печати')
assert get('К2', 'Замечания устранены') == D(7) and get('К2', 'Проверил') == 'Анна' and get('К2', 'П1') == 1
# The Technical expert however written is «ТЭ», a reviewer like the others, dated by the letter of ТЭ when the row has no date.
assert (get('К3', 'Проверил'), get('К3', 'Дата проверки')) == ('ТЭ', D(6)), (get('К3', 'Проверил'), get('К3', 'Дата проверки'))
status = {c: mv.cell(r, COL['Статус комплекта']).value for c, r in codes.items()}
# «Примечание» is a comment in RD by default: Д1 with a note is ready, «Видов замечаний» counts the flags only.
assert status == {'К1': 'готово к ВПР', 'К2': 'готово к ВПР', 'К3': 'замечания', 'Т1': 'готово к ВПР', 'Т2': 'замечания', 'Д1': 'готово к ВПР'}, status
assert mv.cell(codes['Д1'], COL['Видов замечаний']).value == 0 and 'TRIM(' not in m.cell(7, COL['Видов замечаний']).value, m.cell(7, COL['Видов замечаний']).value
assert all(str(m.cell(r, COL[k]).value).startswith('=') for r in range(7, 7 + 6 + 50) for k in ('Сверка с ВЭ РД', 'Статус комплекта', 'Действующая редакция', 'Элемент строки'))
assert '$D$7:$D$62' in m.cell(7, COL['Действующая редакция']).value, m.cell(7, COL['Действующая редакция']).value

# «Элементы»: an element is finished on the day of its last current set, by whoever checked it — ТЭ too.
e, ev = wb['Элементы'], cached['Элементы']
val = lambda x: '' if x is None else x
elements = {ev.cell(r, 1).value: [val(ev.cell(r, c).value) for c in (4, 5, 8, 9, 11, 12, 13, 16, 17, 18)] for r in range(6, 9)}
assert elements == {'1.1': [3, 3, 3, 1, 2, 'замечания', 'не поступал', D(6), 'ТЭ', 1], '1.2': [2, 2, 2, 1, 1, 'замечания', 'замечания', D(7), 'Вера', 0],
                    '1.3': [1, 1, 1, 0, 1, 'готов к ВПР', 'не поступал', D(7), 'Вера', 0]}, elements
assert e['A6'].value.startswith('=IFERROR(INDEX(') and ev['A9'].value in (None, '')

# «СВОД»: days from the first check (working days and the report date), the reviewers under the last day, ТЭ last; no line of ТЭ under the table.
s, sv = wb['СВОД'], cached['СВОД']
assert sv['U2'].value == D(8) and sv['A8'].value == 'Этап 9', (sv['U2'].value, sv['A8'].value)
assert sv['D4'].value == 'Всего комплектов\nс замечаниями: 3' and sv['H4'].value == 'Устранено замечаний\nпо комплектам: 1' and sv['L4'].value == 'Устраняется комплектов\nс замечаниями: 2', (sv['D4'].value, sv['H4'].value, sv['L4'].value)
days = [r for r in range(8, s.max_row + 1) if isinstance(sv.cell(r, 2).value, datetime)]
assert [sv.cell(r, 2).value for r in days] == [D(5), D(6), D(7), D(8)], [sv.cell(r, 2).value for r in days]
assert [[sv.cell(r, c).value for c in (4, 6, 8, 10, 12, 14, 16, 18)] for r in days] == [[3, 6, 0, 2, 3, 4, 0, 2], [3, 6, 1, 4, 2, 2, 1, 2], [3, 6, 3, 6, 0, 0, 2, 2], [3, 6, 3, 6, 0, 0, 0, 0]]
people = {sv.cell(r, 2).value: [sv.cell(r, c).value for c in (8, 10, 16, 18)] for r in range(days[-1] + 1, days[-1] + 5)}
assert people == {'Анна': [0, 2, 0, 0], 'Борис': [0, 1, 0, 0], 'Вера': [2, 2, 0, 0], 'ТЭ': [1, 1, 0, 0]} and sv.cell(days[-1] + 4, 2).value == 'ТЭ', people
assert all(s.row_dimensions[r].hidden for r in range(days[0] + 1, days[0] + 5)) and not s.row_dimensions[days[-1] + 1].hidden
assert all(s.row_dimensions[r].outline_level == 1 for r in range(days[-1] + 1, days[-1] + 5))
assert not any(str(sv.cell(r, 1).value or '').startswith('Замечания ТЭ') for r in range(1, s.max_row + 1)) and sv.cell(days[-1] + 6, 1).value == 'Примечание:'
assert s['J8'].value.startswith('=IF($B8>$U$2,"",COUNTIFS(') and 'Элементы' in s['H8'].value

# «Итоги проверки»: elements finished and sets checked by reviewer; the bold of non-zero values; the alternative table with the acts.
t, tv = wb['Итоги проверки'], cached['Итоги проверки']
assert [tv.cell(4, c).value for c in (3, 5, 7, 9, 11)] == ['Анна', 'Борис', 'Вера', 'ТЭ', 'ВСЕГО']
assert [tv.cell(6, c).value for c in range(3, 13)] == [0, 2, 0, 1, 2, 2, 1, 1, 3, 6], [tv.cell(6, c).value for c in range(3, 13)]
assert [tv.cell(9, c).value for c in range(3, 13)] == [0, 0, 0, 0, 2, 2, 0, 0, 2, 2] and t.cell(9, 7).font.b and not t.cell(9, 3).font.b
alt = next(r for r in range(1, t.max_row + 1) if tv.cell(r, 2).value == 'Борис' and r > 12)
assert [tv.cell(alt, c).value for c in (3, 4, 5, 6)] == [1, 0, 0, 1], [tv.cell(alt, c).value for c in (3, 4, 5, 6)]
assert '$A$1:$S$' in s.print_area and '$A$1:$L$10' in t.print_area, (s.print_area, t.print_area)

# «Статистика»: sets of the stage list, presented, checked, statuses.
st = cached['Статистика']
stats = {st.cell(r, 1).value: st.cell(r, 2).value for r in range(6, st.max_row + 1) if st.cell(r, 1).value}
assert (stats['В ведомости элементов РД (ВЭ РД)'], stats['Представлено на проверку (действующие редакции из ВЭ РД)'], stats['Проверено по состоянию на дату отчёта (как в «СВОД»)'], stats['Готово к ВПР'], stats['Замечания ТЭ («Проверил: ТЭ»)']) == (6, 6, 6, 4, 1), stats
assert 'Примечание (текст в столбце «Примечание»)' not in stats, 'a comment is not a kind of remark'
assert_sums(target)

# --notes-as-remarks: the note is a kind of remark, as in PD.
noted = form(B / 'Примечание.xlsx', [row('К1', '1.1', who='Анна', date=D(5))])
assert 'error' not in build(noted, vera, switches=['--notes-as-remarks'])
nm, nv = load_workbook(noted)[MAIN], load_workbook(noted, data_only=True)[MAIN]
d1 = next(r for r in range(7, nm.max_row + 1) if nm.cell(r, COL['Шифр РД']).value == 'Д1')
assert (nv.cell(d1, COL['Статус комплекта']).value, nv.cell(d1, COL['Видов замечаний']).value) == ('замечания', 1) and 'TRIM($' in nm.cell(7, COL['Видов замечаний']).value
nst = load_workbook(noted, data_only=True)['Статистика']
assert any(nst.cell(r, 1).value == 'Примечание (текст в столбце «Примечание»)' and nst.cell(r, 2).value == 1 for r in range(6, nst.max_row + 1))
assert load_workbook(noted, data_only=True)['СВОД']['D4'].value == 'Всего комплектов\nс замечаниями: 2'
assert_sums(noted)

# The same files again change nothing; a build without files recalculates the reports from the book.
before = {(n, c.coordinate): c.value for n in (MAIN, 'СВОД', 'Итоги проверки', 'Элементы') for r in load_workbook(target, data_only=True)[n].iter_rows() for c in r}
again = build(target, boris, vera)
after = {(n, c.coordinate): c.value for n in (MAIN, 'СВОД', 'Итоги проверки', 'Элементы') for r in load_workbook(target, data_only=True)[n].iter_rows() for c in r}
assert again['Documents'] == 6 and before == after, [k for k in before if before[k] != after.get(k)][:10]
alone = build(target)
assert alone['Files'] == 0 and alone['Documents'] == 6 and alone['Boxes'] == 3, alone

# A form with prepared formula rows: the ranges of the formulas, of other sheets and of the validation follow the data down.
def own_sheet(wb):
    wb.create_sheet('Мой лист')['A1'] = "=COUNTIF('Проверка комплектов'!$D$7:$D$12,\"?*\")"
grown = form(B / 'Рост.xlsx', [row('К1', '1.1', who='Анна', date=D(5))], formula_rows=6, extra=own_sheet)
many = form(B / 'Много.xlsx', [row('Х%d' % i, '1.2' if i == 0 else None, who='Борис', date=D(6)) for i in range(60)])
result = build(grown, many)
assert result['Documents'] == 61, result
assert_sums(grown)

# Letters of ТЭ months before the checks add their own days only, not all the working days from them.
early = form(B / 'Письма ТЭ.xlsx', [row('К1', '1.1', who='Анна', date=D(5)), row('К3', who='ТЭ', ПР=1)])
w = load_workbook(early); w['Ведомость'].cell(9, 11, datetime(2026, 4, 8)); w.save(early)
assert 'error' not in build(early)
sv = load_workbook(early, data_only=True)['СВОД']
assert [sv.cell(r, 2).value for r in range(8, sv.max_row + 1) if isinstance(sv.cell(r, 2).value, datetime)] == [datetime(2026, 4, 8), D(5), D(6), D(7), D(8)]
assert_sums(early)
g = load_workbook(grown)
assert g['Мой лист']['A1'].value == "=COUNTIF('Проверка комплектов'!$D$7:$D$117,\"?*\")", g['Мой лист']['A1'].value
assert '$D$7:$D$117' in g[MAIN].cell(7, COL['Действующая редакция']).value
sets_not_in_stage = [load_workbook(grown, data_only=True)[MAIN].cell(r, COL['Сверка с ВЭ РД']).value for r in range(8, 68)]
assert set(sets_not_in_stage) == {'нет в ВЭ РД'}, set(sets_not_in_stage)

# The row of column numbers moved among the data by a sort: skipped in a reviewer's copy, the consolidated book is not touched.
def numbers_at_end(path):
    wb = load_workbook(path); m = wb[MAIN]; values = [c.value for c in m[6]]; m.delete_rows(6); last = max(r for r in range(1, m.max_row + 1) if m.cell(r, 4).value) + 1
    for c, v in enumerate(values, 1): m.cell(last, c, v)
    wb.save(path); return path
copy = numbers_at_end(form(B / 'Копия с номерами.xlsx', [row('Т1', '1.2', who='Борис', date=D(6))]))
fine = form(B / 'Сводная 2.xlsx', [row('К1', '1.1', who='Анна', date=D(5))])
skipped = build(fine, copy)
assert 'error' not in skipped and skipped['Documents'] == 2 and any('номерами столбцов' in p for p in skipped['ProblemList']), skipped
sorted_target = numbers_at_end(form(B / 'Сводная 3.xlsx', [row('К1', '1.1', who='Анна', date=D(5))]))
before = Path(sorted_target).read_bytes()
refused = build(sorted_target, boris)
assert 'error' in refused and 'номерами столбцов' in refused['error'] and Path(sorted_target).read_bytes() == before, refused

# A book without «Ведомость» is not an RD form: nothing is written.
bare = form(B / 'Без ведомости.xlsx', [row('К1', '1.1', who='Анна', date=D(5))], stage=False)
failed = build(bare, boris)
assert 'error' in failed and 'Ведомость' in failed['error'], failed
# «Приоритет проверки»: 1 within the last 7 working days before «Дата завершения проверки» (and after it), 2 before them; against today, as TODAY().
def workday(day, n):
    while n:
        day -= timedelta(days=1)
        if day.weekday() < 5: n -= 1
    return day
today = date.today()
dt = lambda d: datetime(d.year, d.month, d.day)
due = {'К1': today + timedelta(days=30), 'К2': today + timedelta(days=3), 'К3': today - timedelta(days=10), 'Т1': today + timedelta(days=30)}
expected = {k: 1 if today >= workday(d, 7) else 2 for k, d in due.items()}
assert expected == {'К1': 2, 'К2': 1, 'К3': 1, 'Т1': 2}, expected
DUE = columns(HEADS_DUE)
deadline = form(B / 'Приоритет.xlsx', [row('К1', '1.1', who='Анна', date=D(5), due=dt(due['К1'])), row('К2', due=dt(due['К2'])), row('К3', due=dt(due['К3'])),
                                       row('Т1', '1.2', due=due['Т1'].strftime('%d.%m.%Y')), row('Т2')], heads=HEADS_DUE)
w = load_workbook(deadline); w[MAIN].cell(3, DUE['Дата проверки'], 'дд.мм.гггг'); w.save(deadline)
copy_due = form(B / 'Срок в копии.xlsx', [row('К1', '1.1', who='Анна', date=D(5), due=dt(today)), row('Д1', '1.3', who='Вера', date=D(7), due=dt(due['К2']))], heads=HEADS_DUE)
for _ in range(2): assert 'error' not in build(deadline, copy_due)
pm, pv = load_workbook(deadline)[MAIN], load_workbook(deadline, data_only=True)[MAIN]
rows_due = {pm.cell(r, DUE['Шифр РД']).value: r for r in range(7, pm.max_row + 1) if pm.cell(r, DUE['Шифр РД']).value}
priority = {k: pv.cell(r, DUE['Приоритет проверки ']).value for k, r in rows_due.items()}
assert priority == {**expected, 'Т2': None, 'Д1': 1}, priority
# The deadline of the consolidated book stays; a new set takes it from the file; a date typed as text becomes a date.
assert pv.cell(rows_due['К1'], DUE['Дата завершения проверки']).value == dt(due['К1']) and pv.cell(rows_due['Т1'], DUE['Дата завершения проверки']).value == dt(due['Т1'])
assert pv.cell(rows_due['Д1'], DUE['Дата завершения проверки']).value == dt(due['К2'])
formula = pm.cell(7, DUE['Приоритет проверки ']).value
assert formula == '=IF($D7="","",IF(ISNUMBER($N7),IFERROR(IF(TODAY()>=WORKDAY($N7,-7),1,2),""),""))', formula
assert pv.cell(rows_due['К2'], DUE['Статус комплекта']).value == 'не проверен' and pv.cell(rows_due['К1'], DUE['Статус комплекта']).value == 'готово к ВПР'
rules = [(str(cf.sqref), rule.operator, rule.formula) for cf in pm.conditional_formatting for rule in cf.rules]
assert rules.count(('O7:O%d' % (rows_due['Д1'] + 50), 'equal', ['1'])) == 1, rules
assert (pm.cell(3, DUE['Дата завершения проверки']).value, pm.cell(3, DUE['Приоритет проверки ']).value) == ('дд.мм.гггг', 'автоматически')
assert_sums(deadline)

# «База элементов»: the stage list file updates «Ведомость» (A–AE) at every build; the totals under the list are not copied.
def base_file(path, sets, heading='Статус комплекта'):
    wb = Workbook(); ws = wb.active; ws.title = 'База'
    for c, h in {1: 'Номер элемента', 2: 'Наименование элемента', 3: 'Шифр ПД', 4: 'Шифр РД', 5: 'Наименование документации', 9: heading, 11: 'Дата', 31: '№'}.items(): ws.cell(3, c, h)
    for c in range(1, 32): ws.cell(4, c, c)
    r = 5
    for el, name, items in sets:
        for k, (code, te) in enumerate(items):
            if k == 0: ws.cell(r, 1, el); ws.cell(r, 2, name)
            ws.cell(r, 4, code); ws.cell(r, 5, 'Комплект ' + code); ws.cell(r, 9, te); ws.cell(r, 31, 'В-%d' % r); r += 1
        if len(items) > 1: ws.merge_cells('A%d:A%d' % (r - len(items), r - 1))
    ws.cell(r, 1, '=COUNTA(A5:A%d)' % (r - 1)); ws.cell(r, 4, None)
    wb.save(path); return path
updated = [('1.1', 'Мост', [('К1', 'Да'), ('К2', 'Да'), ('К3', 'Да'), ('К4', 'Да')]), ('1.2', 'Труба', [('Т1', 'Да'), ('Т2', 'Да')])]
stage = base_file(B / 'Этап 9 база.xlsx', updated)
based = form(B / 'С базой.xlsx', [row('К1', '1.1', who='Анна', date=D(5)), row('К3', who='Анна', date=D(5)), row('Д1', '1.3', who='Вера', date=D(7))])
report = build(based, switches=['--rd-base=' + str(stage)])
assert 'error' not in report and any('«Ведомость» обновлена из базы «Этап 9 база.xlsx»: комплектов 6, элементов 2; новых шифров 1, исключено 1' in s for s in report['Log']), report.get('Log', report)
bw, bv = load_workbook(based), load_workbook(based, data_only=True)
ved = bw['Ведомость']
assert [ved.cell(r, 4).value for r in range(7, 14)] == ['К1', 'К2', 'К3', 'К4', 'Т1', 'Т2', None], [ved.cell(r, 4).value for r in range(7, 14)]
assert ved.cell(13, 1).value is None and ved.cell(5, 31).value == '№' and ved.cell(7, 31).value == 'В-5' and 'A7:A10' in {str(x) for x in ved.merged_cells.ranges}
bm = bv[MAIN]; brows = {bm.cell(r, COL['Шифр РД']).value: r for r in range(7, bm.max_row + 1) if bm.cell(r, COL['Шифр РД']).value}
assert bm.cell(brows['К3'], COL['Согласование ТЭ']).value == 'Да' and bm.cell(brows['Д1'], COL['Сверка с ВЭ РД']).value == 'нет в ВЭ РД'
el = bv['Элементы']; assert [el.cell(6, c).value for c in (1, 4, 5)] == ['1.1', 4, 2], [el.cell(6, c).value for c in (1, 4, 5)]
assert_sums(based)
# A base whose columns differ from «Ведомость» is refused; the book does not change.
wrong = base_file(B / 'Чужая база.xlsx', updated, heading='Статус элемента')
before = Path(based).read_bytes()
refused = build(based, switches=['--rd-base=' + str(wrong)])
assert 'error' in refused and 'не совпадает' in refused['error'] and Path(based).read_bytes() == before, refused
print('RD tests passed: merge by set code, new sets after the data, statuses, notes as a comment or a remark, elements, summary, reviewer totals, statistics, PDF, repeated and file-less builds, growing ranges, moved numbers row, missing stage list, priority by the deadline, stage list from the base file')
