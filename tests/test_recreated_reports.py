"""Rebuilt report layouts, reviewer histories and unique global totals; artificial data only."""
from pathlib import Path
from datetime import datetime
import hashlib, subprocess, uuid, math, posixpath, zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter as letter

ROOT = Path(__file__).resolve().parents[1]
B = ROOT / 'tests/.runs' / uuid.uuid4().hex
B.mkdir(parents=True)
REPORTS = {'СВОД', 'Стат. по типам', 'Стат. по зам.', 'Итоги работы'}
MAIN = 'Все загруженные файлы'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
AS_OF = datetime(2026, 10, 4)
EDGES = ('left', 'right', 'top', 'bottom')
EXCLUDED_REVIEWERS = {'Громов', 'Громов Григорий Григорьевич', 'Зорина', 'Дьяков', 'Егоров'}


def assert_report_reviewers(book, included, excluded):
    """Filtering applies to every generated report, not to preserved source rows."""
    for name in REPORTS:
        labels = {cell.value for row in book[name] for cell in row
                  if isinstance(cell.value, str) and cell.data_type != 'f'}
        assert set(included) <= labels, (name, 'missing qualified reviewers', set(included) - labels)
        assert not set(excluded) & labels, (name, 'unqualified reviewers remain', set(excluded) & labels)
    dates = [cell.value for cell in book['СВОД']['B'] if isinstance(cell.value, datetime)]
    # «СВОД»: the first date is in row 8, its reviewers below it.
    first_names = [book['СВОД'].cell(row, 2).value for row in range(9, 9 + len(included))]
    for name in ('Стат. по типам', 'Стат. по зам.'):
        sheet = book[name]
        # The remarks sheet has another table below; date blocks end before its
        # "Вид документации" header. Only the first table contains date cells.
        assert [cell.value for cell in sheet['B'] if isinstance(cell.value, datetime)] == dates
        for index, date in enumerate(dates):
            start = 6 + index * (len(included) + 1)
            assert sheet.cell(start, 2).value == date, (name, date, start)
            names = [sheet.cell(row, 2).value
                     for row in range(start + 1, start + len(included) + 1)]
            assert names == first_names and set(names) == set(included), (name, date, names)
            for offset, reviewer in enumerate(names, 1):
                row = start + offset
                assert sheet.row_dimensions[row].outlineLevel == 1, (name, reviewer, date)
                # Personal statistics stay daily even though every qualified
                # person is now shown on dates when they did no work.
                columns = range(3, 19, 2) if name == 'Стат. по типам' else (5, 7, 9, 11, 13, 15)
                for column in columns:
                    assert_day_formula(sheet.cell(row, column).value, '$B$%d' % start, cumulative=False)


def assert_unframed(sheet, first=1, last=3, width=None):
    """Check every cell, including otherwise empty merged/title cells."""
    for row in sheet.iter_rows(min_row=first, max_row=last, max_col=width or sheet.max_column):
        for cell in row:
            assert not cell.alignment.wrap_text, (sheet.title, cell.coordinate, 'wrapped caption')
            assert all(getattr(cell.border, edge).style in (None, 'none') for edge in EDGES), (
                sheet.title, cell.coordinate, 'border outside table', cell.border)


def sheet_xml(path, name):
    """Resolve sheet parts instead of depending on workbook sheet ordering."""
    with zipfile.ZipFile(path) as archive:
        book = ET.fromstring(archive.read('xl/workbook.xml'))
        rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {rel.get('Id'): rel.get('Target') for rel in rels}
        entry = next(sheet for sheet in book.find('s:sheets', NS) if sheet.get('name') == name)
        target = targets[entry.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
        member = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        return ET.fromstring(archive.read(member))


def serialized_column_width(doc, column):
    matches = [node for node in doc.findall('s:cols/s:col', NS)
               if int(node.get('min')) <= column <= int(node.get('max'))]
    assert len(matches) == 1, (column, matches)
    return float(matches[0].get('width'))


def excel_display_column_width(serialized, maximum_digit_width):
    """OOXML widths include pixel padding; Excel's width dialog does not.

    Keep this inverse conversion independent of the production serializer. The
    caller supplies the Normal font's maximum digit width at 96 dpi; report cell
    fonts do not determine the column-width unit.
    """
    padding = 2 * math.ceil(maximum_digit_width / 4) + 1
    pixels = math.floor(serialized * maximum_digit_width + 0.5)
    return math.floor((pixels - padding) / maximum_digit_width * 100 + 0.5) / 100


def assert_cached(formulas, cached, row, column, expected):
    cell = formulas.cell(row, column)
    actual = cached.cell(row, column).value
    assert cell.data_type == 'f', (formulas.title, cell.coordinate, cell.value)
    assert '#REF!' not in cell.value, cell.value
    if expected is None:
        assert actual in (None, ''), (formulas.title, cell.coordinate, actual, 'future should be blank')
    else:
        assert isinstance(actual, (int, float)) and math.isclose(actual, expected, abs_tol=1e-9), (
            formulas.title, cell.coordinate, actual, expected, cell.value)
    return cell.value


def assert_day_formula(formula, anchor, cumulative):
    assert 'COUNTIFS(' in formula and "'Справка по томам'!" in formula, formula
    assert anchor in formula, formula
    if cumulative:
        assert '\">0\"' in formula and '\"<=\"&' + anchor in formula, formula
    else:
        assert '\"<=\"' not in formula, formula


def report_snapshot(path):
    """Compare semantic styles and values, not regenerated style IDs or ZIP bytes."""
    formulas, cached = load_workbook(path), load_workbook(path, data_only=True)
    result = {}
    for name in sorted(REPORTS):
        sheet = formulas[name]
        cells = []
        for row in sheet:
            for cell in row:
                color = cell.font.color
                cells.append((cell.coordinate, cell.value, cached[name][cell.coordinate].value,
                              cell.font.name, cell.font.sz, cell.font.bold,
                              None if color is None else (color.type, str(color.value)),
                              cell.number_format, cell.alignment.wrap_text,
                              tuple(getattr(cell.border, edge).style for edge in EDGES),
                              cell.fill.patternType, cell.fill.fgColor.type, cell.fill.fgColor.index))
        doc = sheet_xml(path, name)
        result[name] = (cells, sorted(map(str, sheet.merged_cells.ranges)), str(sheet.print_area),
                        doc.find('s:dimension', NS).get('ref'),
                        [(int(col.get('min')), int(col.get('max')), float(col.get('width')))
                         for col in doc.findall('s:cols/s:col', NS)],
                        [(r, d.outlineLevel, d.hidden, d.collapsed, d.height)
                         for r, d in sorted(sheet.row_dimensions.items())])
    formulas.close()
    cached.close()
    return result

def run(output, source, as_of=AS_OF, *switches):
    subprocess.run([str(ROOT / 'Свод_проверки.exe'), '--batch', str(output),
                    as_of.strftime('%Y-%m-%d'), MAIN, str(source), *switches], check=True, timeout=90)


def report_values(path):
    """Cached numbers and texts of the four reports, without formulas."""
    cached = load_workbook(path, data_only=True)
    return {name: {c.coordinate: c.value for row in cached[name] for c in row} for name in REPORTS}

def validate(path):
    w = load_workbook(path)
    assert REPORTS <= set(w.sheetnames), w.sheetnames
    assert 'Свод' not in w.sheetnames
    # Row 4 of «СВОД» holds the volumes with remarks; the table header starts at row 5.
    assert w['СВОД']['D5'].value == 'Всего'
    assert w['СВОД']['H5'].value == 'Проверено'
    assert w['СВОД']['L5'].value == 'Осталось проверить'
    assert w['СВОД']['P5'].value == 'Проверено за сутки'
    assert w['Стат. по зам.']['A1'].value == 'СТАТИСТИКА ПО ЗАМЕЧАНИЯМ'
    assert w['Итоги работы']['A1'].value == 'СВОДНАЯ ТАБЛИЦА ПО ИТОГАМ РАБОТЫ ПРОВЕРЯЮЩИХ'
    assert w['Итоги работы']['C5'].value == 'коробов'
    assert w['Итоги работы']['D5'].value == 'томов'
    for tab in REPORTS:
        s = w[tab]
        assert_unframed(s)
        header = 'A5' if tab == 'СВОД' else 'A4'
        assert s[header].value == 'Этап' and s[header].border.top.style == 'thick', (tab, header)
        assert s.print_area and s.page_setup.orientation == 'landscape'
        assert s['A1'].font.name == 'Times New Roman'
        for row in s:
            for c in row:
                if c.data_type == 'f':
                    assert '#REF!' not in c.value
                    assert 'IF(' in c.value or 'COUNT' in c.value or 'SUM' in c.value
    stats = w['Стат. по зам.']
    assert [stats[a].font.color.rgb for a in ['C6', 'E6', 'I6', 'M6']] == ['FF996600', 'FF00B050', 'FFFF0000', 'FF0070C0']
    assert w['СВОД']['D7'].border.right.style == 'dashed' and w['СВОД']['E7'].border.left.style == 'dashed'
    assert w[MAIN]['M5'].value == 'Номер тома' and w[MAIN]['M7'].number_format == '@'
    assert sum(bool(view.tabSelected) for s in w for view in s.views.sheetView) == 1
    assert any(stats.row_dimensions[r].outlineLevel == 1 for r in stats.row_dimensions)
    return w

def test_layout_migration():
    # Reproduce the user's old workbook with a horizontal report and no required captions.
    source = B / 'old-layout.xlsx'
    w = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx')
    for s in list(w):
        if s.title.upper() == 'СВОД': w.remove(s)
    s = w.create_sheet('Свод')
    s['A1'] = 'Старое оформление'
    s['A7'] = 'Получено коробов'; s['A8'] = 'Томов по акту'
    s['B5'] = datetime(2026,10,2); s['C5'] = datetime(2026,10,4)
    s['B7'] = "='Справка по томам'!B1"; s['B8'] = "='Справка по томам'!B2"
    s['A17'] = 'Примечание: исходные сведения сохраняются'
    if 'Справка по томам' not in w.sheetnames: w.create_sheet('Справка по томам')
    w['Справка по томам']['B1'] = 89; w['Справка по томам']['B2'] = 1975
    w.save(source); digest = hashlib.sha256(source.read_bytes()).hexdigest()
    target = B / 'recreated.xlsx'; run(target, source); w = validate(target)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert w['СВОД']['A1'].value != 'Старое оформление'
    assert any('исходные сведения' in str(c.value) for row in w['СВОД'] for c in row)
    assert w['Справка по томам']['B1'].value == 89 and w['Справка по томам']['B2'].value == 1975

    # Existing report tabs must also be regenerated, including stale formatting and formulas.
    for tab in REPORTS:
        w[tab]['A1'] = 'ПОВРЕЖДЕНО'
        w[tab]['A4'].border = Border()
        w[tab]['C6'].font = Font(color='FF00FF')
    w.save(target); run(target, source); w = validate(target)
    assert all(w[s]['A1'].value != 'ПОВРЕЖДЕНО' for s in REPORTS)
    assert w['Справка по томам']['B2'].value == 1975
    no_layout = B / 'no-layout.xlsx'
    w = load_workbook(source); w.remove(w['Свод']); w.save(no_layout)
    fresh = B / 'new.xlsx'; run(fresh, no_layout); validate(fresh)
    # A custom incomplete template previously entered the silent "keep old report" branch.
    w = load_workbook(source); s = w['Свод']; s['A4']='Этап'; s['B4']='Дата / проверяющий'; w.save(source)
    custom = B / 'custom.xlsx'; run(custom, source); validate(custom)


def artificial_source(path, empty=False):
    """Use real register headers with entirely artificial, deterministic review rows."""
    book = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx')
    # Widths are measured against Normal, not the Times New Roman report cells.
    # Courier New 12 has a deterministic maximum digit width of 10 px at 96 dpi.
    book._named_styles['Normal'].font = Font(name='Courier New', size=12)
    main = book[MAIN]
    main.delete_rows(7, main.max_row)
    for sheet in list(book):
        if sheet.title.upper() == 'СВОД' or sheet.title in REPORTS or sheet.title == 'Справка по томам':
            book.remove(sheet)
    # A source calendar makes the first and last report dates explicit. On a repeat
    # merge these dates are also recoverable from dated register rows plus as-of.
    dates = [AS_OF] if empty else [datetime(2026, 10, day) for day in range(1, 6)]
    old = book.create_sheet('Свод')
    old['A7'], old['A8'] = 'Получено коробов', 'Томов по акту'
    for column, date in enumerate(dates, 2):
        old.cell(5, column, date)
    ref = book.create_sheet('Справка по томам')
    ref['A1'], ref['B1'] = 'Коробов получено', 20
    ref['A2'], ref['B2'] = 'Томов по проекту акта', 100
    ref['A8'] = 'ФИО проверяющего'
    ref['A6'] = 'Примечание: сохранить сведения из исходной справки'
    ref['A42'], ref['A43'] = 'ИТОГО проверено коробов', 'ИТОГО проверено томов'
    if empty:
        records = [('Без проверки', None, None, 11)]
    else:
        ref['A12'] = 'Громов Григорий Григорьевич'  # Not present in any file row.
        records = [
            ('Совместный_Часть1', 'Анна', 1, 11),
            ('Совместный_Часть2', 'Борис', 2, 11),
            ('Совместный-ИУЛ', 'Анна', 1, None),
            ('Первый том', 'Анна', 1, 12),
            ('Второй том', 'Борис', 2, 11),  # Same physical box as the shared volume.
            ('Поздний том', 'Анна', 3, 13),
            ('Будущий том', 'Вера', 5, 14),
            ('Без даты', 'Зорина', None, 15),  # Reviewer exists, but has no dated work.
            ('Без проверяющего', None, None, 16),
            ('Отдельный-ИУЛ', 'Борис', 2, 17),  # A checked file/box, never a volume.
            ('Служебный-ИУЛ', 'Дьяков', 2, 18),  # No completed volume in this reviewer's history.
            ('Незавершенный_Часть1', 'Егоров', 2, 19),
            ('Незавершенный_Часть2', None, None, 19),  # The shared volume is incomplete.
        ]
    for row, (filename, reviewer, day, box) in enumerate(records, 7):
        for column, value in {1: row - 6, 2: 'Искусственный раздел', 3: filename,
                              4: 'pdf', 5: '1 МБ', 6: '01.10.2026', 7: '10:00',
                              8: 'FEED%04d' % row, 9: reviewer,
                              10: None if day is None else datetime(2026, 10, day),
                              11: 'ПД', 12: box}.items():
            main.cell(row, column, value)
    # Deliberately hostile legacy caption formatting must not leak into outputs.
    heavy = Side(style='thick', color='FF000000')
    for row in old.iter_rows(min_row=1, max_row=3, max_col=19):
        for cell in row:
            cell.border = Border(left=heavy, right=heavy, top=heavy, bottom=heavy)
            cell.alignment = Alignment(wrap_text=True)
    book.save(path)
    book.close()
    return dates


def assert_remarks_row(summary, last):
    """The volumes with remarks are in row 4, right above the «СВОД» table; the table ends at the last row of the sheet."""
    text = str(summary.cell(4, 4).value)
    assert summary.max_row == last and text.startswith('="Всего томов"&CHAR(10)&"с замечаниями: "'), (last, summary.max_row, text)
    assert summary.cell(5, 1).value == 'Этап' and summary.cell(8, 1).value == 'Этап 1'


def assert_matrix_layout(path, sheet, reviewers, dates, maximum_digit_width=10):
    """Two columns per reviewer plus a real, final two-column ВСЕГО group."""
    end = 4 + 2 * len(reviewers)
    total = end - 1
    last = 6 + len(dates)
    headers = {sheet.cell(4, column).value: column for column in range(3, end + 1, 2)}
    assert set(headers) == set(reviewers) | {'ВСЕГО'}, headers
    assert headers['ВСЕГО'] == total, headers
    assert sheet['B6'].value == 'Всего'
    assert len(headers) == len(reviewers) + 1
    normal = sheet.parent._named_styles['Normal'].font
    assert normal.name == 'Courier New' and normal.sz == 12, normal
    assert sheet['C5'].font.name == 'Times New Roman', sheet['C5'].font
    merges = set(map(str, sheet.merged_cells.ranges))
    # The title is not merged: it starts in A1, left aligned, and runs over the empty cells to the right.
    assert not any(m.startswith('A1:') for m in merges), merges
    assert sheet['A1'].value == 'СВОДНАЯ ТАБЛИЦА ПО ИТОГАМ РАБОТЫ ПРОВЕРЯЮЩИХ' and sheet['A1'].alignment.horizontal == 'left'
    assert 'A4:A5' in merges and 'B4:B5' in merges
    for column in range(3, end + 1, 2):
        assert '%s4:%s4' % (letter(column), letter(column + 1)) in merges
        assert sheet.cell(5, column).value == 'коробов'
        assert sheet.cell(5, column + 1).value == 'томов'
        assert sheet.cell(4, column).font.bold
        for row in range(5, last + 1):
            assert sheet.cell(row, column).border.right.style == 'thin'
            assert sheet.cell(row, column + 1).border.left.style == 'thin'
    # Inspect actual OOXML col intervals: openpyxl does not expand grouped widths.
    doc = sheet_xml(path, sheet.title)
    widths = [serialized_column_width(doc, column) for column in range(3, end + 1)]
    assert len(set(widths)) == 1, widths
    target_chars = 8.9
    padding = 2 * math.ceil(maximum_digit_width / 4) + 1
    expected_pixels = math.floor(target_chars * maximum_digit_width + padding + 0.5)
    expected_raw = math.floor(expected_pixels / maximum_digit_width * 256) / 256
    for column, raw in enumerate(widths, 3):
        display = excel_display_column_width(raw, maximum_digit_width)
        assert raw == expected_raw, (column, raw, expected_raw, 'serialized OOXML width')
        assert display == target_chars, (column, raw, display, 'calculated Excel UI width')
    assert expected_raw == 9.59765625 and expected_pixels == 96
    assert doc.find('s:dimension', NS).get('ref').split(':')[-1].rstrip('0123456789') == letter(max(8, end))
    areas = str(sheet.print_area)
    assert '$A$1:$%s$%d' % (letter(end), last) in areas, areas
    # The matrix must have no nonempty columns after ВСЕГО, even when a wider
    # alternative table below requires eight sheet columns (zero-reviewer case).
    for row in sheet.iter_rows(min_row=4, max_row=last, min_col=end + 1,
                               max_col=max(end + 1, sheet.max_column)):
        assert all(cell.value is None for cell in row), row
    for row in range(4, last + 1):
        assert sheet.cell(row, 1).border.left.style == 'thick'
        assert sheet.cell(row, end).border.right.style == 'thick'
        for column in range(3, end, 2):
            assert sheet.cell(row, column).border.left.style == 'medium'
            assert sheet.cell(row, column - 1).border.right.style == 'medium'
    for column in range(1, end + 1):
        assert sheet.cell(4, column).border.top.style == 'thick'
        assert sheet.cell(last, column).border.bottom.style == 'thick'
    assert_unframed(sheet)
    alternative = next(cell.row for row in sheet for cell in row
                       if cell.value == 'Альтернативное представление по проверяющим')
    assert alternative == last + 3
    assert_unframed(sheet, alternative, alternative, 8)
    assert [sheet.cell(7 + index, 2).value for index in range(len(dates))] == dates
    return headers


def test_reviewer_histories_and_global_totals():
    source, target = B / 'reviewer-history-source.xlsx', B / 'reviewer-history.xlsx'
    dates = artificial_source(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    run(target, source)
    book, cached = load_workbook(target), load_workbook(target, data_only=True)
    reviewers = {'Анна', 'Борис', 'Вера'}
    summary, summary_values = book['СВОД'], cached['СВОД']
    work, work_values = book['Итоги работы'], cached['Итоги работы']
    for name in REPORTS:
        assert_unframed(book[name])
    header = assert_matrix_layout(target, work, reviewers, dates)
    assert_report_reviewers(book, reviewers, EXCLUDED_REVIEWERS)
    # October 4 is idle for everyone: these personal statistics must remain
    # daily zeros while the main report's checked columns stay cumulative.
    idle_start = 6 + dates.index(AS_OF) * (len(reviewers) + 1)
    for row in range(idle_start + 1, idle_start + len(reviewers) + 1):
        for column in (5, 7, 9, 11, 13, 15):
            assert_cached(book['Стат. по зам.'], cached['Стат. по зам.'], row, column, 0)
        for column in range(3, 23, 2):
            assert_cached(book['Стат. по типам'], cached['Стат. по типам'], row, column, 0)
    # Expected unique completions: box 11 first counted for Anna on October 1;
    # its two-part/ИУЛ volume counted once for Boris on October 2. A standalone
    # ИУЛ is never a volume. The two excluded dated-file reviewers add boxes
    # to global totals, but have no completed volumes and no personal columns.
    daily = {
        1: {'Анна': (2, 1)},
        2: {'Борис': (1, 2)},
        3: {'Анна': (1, 1)},
        4: {},
    }
    global_daily = {1: (2, 1), 2: (3, 2), 3: (1, 1), 4: (0, 0)}
    cumulative = {name: [0, 0] for name in reviewers}
    global_cumulative = [0, 0]
    assert_remarks_row(summary, 7 + len(dates) * (len(reviewers) + 1))
    assert summary.max_column == 19
    assert sheet_xml(target, 'СВОД').find('s:dimension', NS).get('ref') == 'A1:S27'
    seen = []
    for index, date in enumerate(dates):
        row = 8 + index * (len(reviewers) + 1)
        assert summary.cell(row, 2).value == date
        assert summary.cell(row, 2).font.bold
        names = [summary.cell(r, 2).value for r in range(row + 1, row + len(reviewers) + 1)]
        assert len(names) == len(set(names)) and set(names) == reviewers, (date, names)
        seen.append(names)
        future = date > AS_OF
        if not future:
            for name in reviewers:
                for kind in range(2):
                    cumulative[name][kind] += daily[date.day].get(name, (0, 0))[kind]
            for kind in range(2):
                global_cumulative[kind] += global_daily[date.day][kind]
        for offset, name in enumerate([None] + names):
            r = row + offset
            anchor = '$B%d' % row if name is None else '$B$%d' % row
            # Only the report date is expanded: earlier and future dates stay collapsed.
            if name is None:
                assert bool(summary.row_dimensions[r].collapsed) == (date != AS_OF), (date, 'collapsed mark of the date row')
            else:
                assert bool(summary.row_dimensions[r].hidden) == (date != AS_OF), (date, name, 'expanded report date')
            if name is not None:
                assert summary.row_dimensions[r].outlineLevel == 1
                for column in list(range(4, 8)) + list(range(12, 16)):
                    cell = summary.cell(r, column)
                    assert cell.value in (None, '') and cell.data_type != 'f', (cell.coordinate, cell.value)
                    assert summary_values.cell(r, column).value in (None, '')
            for kind, denominator in enumerate((20, 100)):
                done = global_cumulative[kind] if name is None else cumulative[name][kind]
                today = None if future else (global_daily[date.day][kind] if name is None
                                             else daily[date.day].get(name, (0, 0))[kind])
                formula = assert_cached(summary, summary_values, r, 8 + kind * 2, None if future else done)
                assert_day_formula(formula, anchor, cumulative=True)
                assert_cached(summary, summary_values, r, 9 + kind * 2,
                              None if future else done / denominator * 100)
                formula = assert_cached(summary, summary_values, r, 16 + kind * 2, today)
                assert_day_formula(formula, anchor, cumulative=False)
                assert_cached(summary, summary_values, r, 17 + kind * 2,
                              None if future else today / denominator * 100)
                if name is None:
                    assert_cached(summary, summary_values, r, 4 + kind * 2, None if future else denominator)
                    assert_cached(summary, summary_values, r, 5 + kind * 2, None if future else 100)
                    assert_cached(summary, summary_values, r, 12 + kind * 2,
                                  None if future else denominator - done)
                    assert_cached(summary, summary_values, r, 13 + kind * 2,
                                  None if future else (denominator - done) / denominator * 100)
        matrix_row = 7 + index
        for name, column in header.items():
            for kind in range(2):
                value = None if future else (global_daily[date.day][kind] if name == 'ВСЕГО'
                                              else daily[date.day].get(name, (0, 0))[kind])
                formula = assert_cached(work, work_values, matrix_row, column + kind, value)
                assert_day_formula(formula, '$B%d' % matrix_row, cumulative=False)
                if name == 'ВСЕГО':
                    assert all('"' + who + '"' not in formula for who in reviewers), formula
    assert all(names == seen[0] for names in seen), 'Reviewer order must be stable under every date'
    for name, column in header.items():
        expected = global_cumulative if name == 'ВСЕГО' else cumulative[name]
        for kind in range(2):
            formula = assert_cached(work, work_values, 6, column + kind, expected[kind])
            assert 'COUNTIFS(' in formula and '\"<=\"&' in formula, formula
            # Non-zero values are bold, zeros are not.
            assert bool(work.cell(6, column + kind).font.bold) == (expected[kind] != 0), (name, kind, expected[kind])
            if name == 'ВСЕГО':
                assert all('"' + who + '"' not in formula for who in reviewers), formula
    assert tuple(global_cumulative) == (6, 4)
    assert sum(counts[0] for counts in cumulative.values()) == 4, (
        'Global box count includes the two excluded reviewers; summing displayed columns is wrong')
    assert tuple(cumulative['Анна']) == (3, 2)
    assert tuple(cumulative['Борис']) == (1, 2)
    # Source received totals remain independent of the reviewed global totals.
    assert cached['Справка по томам']['B1'].value == 20
    assert cached['Справка по томам']['B2'].value == 100
    assert cached['Справка по томам']['A12'].value == 'Громов Григорий Григорьевич'
    assert cached['Справка по томам']['A6'].value == 'Примечание: сохранить сведения из исходной справки'
    book.close()
    cached.close()
    before = report_snapshot(target)
    run(target, source, AS_OF, '--recalculate-history')
    assert report_snapshot(target) == before, 'Repeat merge changed formulas, caches, grouping, dimensions or styles'
    # A plain repeat build gives the same numbers; every day, past ones too, stays a formula of the current data.
    values = report_values(target)
    run(target, source)
    assert report_values(target) == values, 'A repeat build changed a number'
    summary = load_workbook(target)['СВОД']
    rows = {summary.cell(r, 2).value: r for r in range(8, summary.max_row + 1) if isinstance(summary.cell(r, 2).value, datetime)}
    past, today = rows[min(rows)], rows[AS_OF]
    assert str(summary.cell(past, 10).value).startswith('=') and str(summary.cell(today, 10).value).startswith('='), (summary.cell(past, 10).value, summary.cell(today, 10).value)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def test_no_reviewers():
    source, target = B / 'no-reviewers-source.xlsx', B / 'no-reviewers.xlsx'
    dates = artificial_source(source, empty=True)
    run(target, source)
    book, cached = load_workbook(target), load_workbook(target, data_only=True)
    headers = assert_matrix_layout(target, book['Итоги работы'], set(), dates)
    assert headers == {'ВСЕГО': 3}
    for row in (6, 7):
        for column in (3, 4):
            assert_cached(book['Итоги работы'], cached['Итоги работы'], row, column, 0)
    summary = book['СВОД']
    assert_remarks_row(summary, 8)
    assert summary.max_column == 19
    assert summary['B8'].value == AS_OF
    assert not summary.row_dimensions[8].outlineLevel
    assert not any(d.outlineLevel for d in summary.row_dimensions.values())
    for column in (8, 9, 10, 11, 16, 17, 18, 19):
        assert_cached(summary, cached['СВОД'], 8, column, 0)
    assert_cached(summary, cached['СВОД'], 8, 4, 20)
    assert_cached(summary, cached['СВОД'], 8, 6, 100)
    assert_cached(summary, cached['СВОД'], 8, 12, 20)
    assert_cached(summary, cached['СВОД'], 8, 14, 100)
    for name in REPORTS:
        assert_unframed(book[name])
    book.close()
    cached.close()
    before = report_snapshot(target)
    run(target, source, AS_OF, '--recalculate-history')
    assert report_snapshot(target) == before, 'Zero-reviewer repeat merge changed reports'


def test_same_surname_reference_reviewer():
    """A reference-only person must not absorb an active namesake's statistics."""
    source, target = B / 'namesakes-source.xlsx', B / 'namesakes.xlsx'
    dates = artificial_source(source)
    book = load_workbook(source)
    active, idle = 'Иванов Иван Иванович', 'Иванов Игорь Игоревич'
    book['Справка по томам']['A12'] = idle
    for row in book[MAIN]:
        if row[8].value == 'Анна':
            row[8].value = active
    book.save(source)
    book.close()
    run(target, source)
    book, cached = load_workbook(target), load_workbook(target, data_only=True)
    reviewers = {active, 'Борис', 'Вера'}
    headers = assert_matrix_layout(target, book['Итоги работы'], reviewers, dates)
    assert 'Иванов' not in headers and idle not in headers
    assert_report_reviewers(book, reviewers, EXCLUDED_REVIEWERS | {idle, 'Иванов'})
    assert cached['Справка по томам']['A12'].value == idle
    for name, expected in ((active, (3, 2)),):
        for kind in range(2):
            formula = assert_cached(book['Итоги работы'], cached['Итоги работы'],
                                    6, headers[name] + kind, expected[kind])
            assert '"' + name + '"' in formula, formula
    summary = book['СВОД']
    for index, date in enumerate(dates):
        start = 8 + index * (len(reviewers) + 1)
        rows = {summary.cell(row, 2).value: row for row in range(start + 1, start + len(reviewers) + 1)}
        assert set(rows) == reviewers, rows
        assert idle not in rows
        for column in (8, 10, 16, 18):
            formula = summary.cell(rows[active], column).value
            assert '"' + active + '"' in formula and '"' + idle + '"' not in formula, formula
    book.close()
    cached.close()


def test_lifetime_reviewers_do_not_depend_on_as_of():
    """Future completed work qualifies a reviewer before that work is reportable."""
    source = B / 'lifetime-reviewers-source.xlsx'
    dates = artificial_source(source)
    reviewers = {'Анна', 'Борис', 'Вера'}
    for as_of in (datetime(2026, 10, 1), datetime(2026, 10, 6)):
        target = B / ('lifetime-reviewers-%s.xlsx' % as_of.strftime('%Y%m%d'))
        run(target, source, as_of)
        book, cached = load_workbook(target), load_workbook(target, data_only=True)
        # A date of the old «Свод» after the report date without checks (October 4) is not a report day.
        expected_dates = sorted(d for d in set(dates + [as_of]) if d <= as_of or d.day != 4)
        headers = assert_matrix_layout(target, book['Итоги работы'], reviewers, expected_dates)
        assert_report_reviewers(book, reviewers, EXCLUDED_REVIEWERS)
        summary = book['СВОД']
        assert_remarks_row(summary, 7 + len(expected_dates) * (len(reviewers) + 1))
        for index, date in enumerate(expected_dates):
            start = 8 + index * (len(reviewers) + 1)
            assert summary.cell(start, 2).value == date
            rows = {summary.cell(row, 2).value: row
                    for row in range(start + 1, start + len(reviewers) + 1)}
            assert set(rows) == reviewers, (as_of, date, rows)
            for name in reviewers:
                if date > as_of:
                    for column in (8, 9, 10, 11, 16, 17, 18, 19):
                        assert_cached(summary, cached['СВОД'], rows[name], column, None)
        # Vera has exactly one volume, completed on October 5. Earlier reports
        # still show her lifetime-qualified column with zero as-of totals.
        for kind in range(2):
            assert_cached(book['Итоги работы'], cached['Итоги работы'], 6,
                          headers['Вера'] + kind, 0 if as_of.day == 1 else 1)
        work = book['Итоги работы']
        alternative = next(cell.row for cell in work['A']
                           if cell.value == 'Альтернативное представление по проверяющим')
        for row in range(alternative + 2, work.max_row + 1):
            date = work.cell(row, 2).value
            if isinstance(date, datetime) and date > as_of:
                for column in range(3, 9):
                    formula = assert_cached(work, cached['Итоги работы'], row, column, None)
                    assert formula.startswith('=IF($B%d>' % row), formula
        book.close()
        cached.close()


def test_only_unqualified_reviewers():
    """Named reviewed files alone must not create a reviewer table column or row."""
    source, target = B / 'unqualified-source.xlsx', B / 'unqualified.xlsx'
    dates = artificial_source(source)
    book = load_workbook(source)
    main = book[MAIN]
    keep = {'Без даты', 'Без проверяющего', 'Служебный-ИУЛ',
            'Незавершенный_Часть1', 'Незавершенный_Часть2'}
    for row in range(main.max_row, 6, -1):
        if main.cell(row, 3).value not in keep:
            main.delete_rows(row)
    book.save(source)
    book.close()
    run(target, source)
    # October 5 of the old «Свод» is after the report date and has no checks of the kept rows.
    dates = [d for d in dates if d <= AS_OF]
    book, cached = load_workbook(target), load_workbook(target, data_only=True)
    headers = assert_matrix_layout(target, book['Итоги работы'], set(), dates)
    assert headers == {'ВСЕГО': 3}
    assert_report_reviewers(book, set(), EXCLUDED_REVIEWERS | {'Анна', 'Борис', 'Вера'})
    assert_cached(book['Итоги работы'], cached['Итоги работы'], 6, 3, 2)
    assert_cached(book['Итоги работы'], cached['Итоги работы'], 6, 4, 0)
    summary = book['СВОД']
    assert_remarks_row(summary, 7 + len(dates))
    assert [summary.cell(row, 2).value for row in range(8, summary.max_row + 1)] == dates
    assert not any(dim.outlineLevel for dim in summary.row_dimensions.values())
    as_of_row = 8 + dates.index(AS_OF)
    assert_cached(summary, cached['СВОД'], as_of_row, 8, 2)
    assert_cached(summary, cached['СВОД'], as_of_row, 10, 0)
    assert cached['Справка по томам']['A12'].value == 'Громов Григорий Григорьевич'
    book.close()
    cached.close()


if __name__ == '__main__':
    test_layout_migration()
    test_reviewer_histories_and_global_totals()
    test_no_reviewers()
    test_same_surname_reference_reviewer()
    test_lifetime_reviewers_do_not_depend_on_as_of()
    test_only_unqualified_reviewers()
    print('Report recreation passed: old/absent/custom layouts; all-date cumulative reviewer rows; '
          'daily and unique global cached totals; future dates; lifetime completed-volume filtering; '
          'reference-only/undated/IUL-only/incomplete/zero reviewers; namesakes; '
          'widths, merged groups, borders, unwrapped captions and repeat merge')
    print(B)
