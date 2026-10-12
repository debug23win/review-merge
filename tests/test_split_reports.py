"""Report migration, shared formulas, automatic column and repeat merge, synthetic data only."""
from pathlib import Path
from datetime import datetime
from copy import copy
import subprocess,uuid,zipfile
import xml.etree.ElementTree as E
from openpyxl import load_workbook
from openpyxl.styles import Font,Border,Side

ROOT=Path(__file__).resolve().parents[1]
B=ROOT/'tests/.runs'/uuid.uuid4().hex;B.mkdir(parents=True)
source=B/'combined.xlsx';target=B/'split.xlsx';seed=B/'seed.xlsx'
subprocess.run([str(ROOT/'Свод_проверки.exe'),'--batch',str(seed),'2026-10-04','Все загруженные файлы',str(ROOT/'tests/fixtures/Анна.xlsx')],check=True,timeout=90)
w=load_workbook(seed)
w['Все загруженные файлы'].delete_cols(13)
for tab in list(w):
 if tab.title.upper()=='СВОД':w.remove(tab)
s=w.create_sheet('Свод')
s['A1']='Сводная таблица';s['A2']='Этап 1';s['A3']='Пример'
s['A4']='Этап';s['B4']='Дата / проверяющий'
for c in range(4,20):s.cell(4,c,'Количество');s.cell(5,c,'%' if c%2 else 'Шт.')
s['B6']=datetime(2026,10,2);s['B7']='Анна';s['B8']=datetime(2026,10,4);s['A9']='Примечание:'
s['D6']="=COUNTIF('Все загруженные файлы'!M7:M11,1)";s['F6']=100
s['A13']='КОЛИЧЕСТВО ЗАМЕЧАНИЙ ПО ТИПАМ'
s['A14']='Этап';s['B14']='Дата / проверяющий'
for c in range(3,23):s.cell(14,c,'Тип');s.cell(15,c,'Шт.' if c%2 else '%')
for r in [16,18]:s.cell(r,2,datetime(2026,10,2 if r==16 else 4));s.cell(r,3,2);s.cell(r,4,'=C'+str(r)+'/100')
s['B17']='Анна';s['C17']=1;s['D17']='=C17/100'
s['A20']='СТАТИСТИКА ПО ОСНОВНОМУ ЛИСТУ';s['A21']='Этап';s['B21']='Дата / проверяющий'
for c in range(3,17):s.cell(21,c,'Показатель');s.cell(22,c,'Шт.' if c%2 else '%')
for r in [23,25]:s.cell(r,2,datetime(2026,10,2 if r==23 else 4));s.cell(r,3,5);s.cell(r,4,'=C'+str(r)+'/5')
s['B24']='Анна';s['A27']='Этап';s['B27']='Вид документации'
for c in range(3,7):s.cell(27,c,'Файлов')
s['B28']='ПД';s['C28']=5
s['A32']='ПРОВЕРЕНО ПО ПРОВЕРЯЮЩИМ';s['A33']='Этап';s['B33']='Дата';s['E33']='=$B$44';s['G33']='Всего'
s['E34']='=$C$44';s['E35']='=$H$44';s['G34']='=SUM(E34)';s['G35']='=SUM(E35)'
for r,day,alt in [(36,2,45),(38,4,46)]:
 s.cell(r,2,datetime(2026,10,day));s.cell(r,5,'=$C$'+str(alt));s.cell(r+1,5,'=$H$'+str(alt));s.cell(r,7,'=SUM(E'+str(r)+')');s.cell(r+1,7,'=SUM(E'+str(r+1)+')')
s['A42']='АЛЬТЕРНАТИВНОЕ ПРЕДСТАВЛЕНИЕ ПО ПРОВЕРЯЮЩИМ'
for c,label in enumerate(['Этап','Проверяющий / дата','Проверено томов','Замечания','За день','Файлы','Без даты','Проверено коробов'],1):s.cell(43,c,label)
s['B44']='Анна';s['C44']='=SUM(C45:C46)';s['H44']='=SUM(H45:H46)'
for r,day,toms in [(45,2,2),(46,4,1)]:s.cell(r,2,datetime(2026,10,day));s.cell(r,3,toms);s.cell(r,8,1)
thin=Side(style='thin',color='000000')
for row in s.iter_rows(max_row=46,max_col=22):
 for cell in row:cell.font=Font(name='Times New Roman',size=12);cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
s.merge_cells('E33:F33');s.merge_cells('G33:H33');w.save(source)
# Real Excel shared-formula storage: the followers have no formula text of their own.
n='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
with zipfile.ZipFile(source) as z:files={a:z.read(a) for a in z.namelist()}
sheet=next(a for a in files if a.startswith('xl/worksheets/sheet') and 'ПРОВЕРЕНО ПО ПРОВЕРЯЮЩИМ' in ''.join(E.fromstring(files[a]).itertext()));doc=E.fromstring(files[sheet]);cells={c.get('r'):c for c in doc.iter(n+'c')}
for a in ['D16','D17','D18']:
 f=cells[a].find(n+'f');f.set('t','shared');f.set('si','7')
 if a=='D16':f.set('ref','D16:D18')
 else:f.text=None
files[sheet]=E.tostring(doc,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(source,'w',zipfile.ZIP_DEFLATED) as z:
 for a,data in files.items():z.writestr(a,data)
def run():subprocess.run([str(ROOT/'Свод_проверки.exe'),'--batch',str(target),'2026-10-04','Все загруженные файлы',str(source)],check=True,timeout=90)
run();w=load_workbook(target)
assert {'СВОД','Стат. по типам','Стат. по зам.','Итоги работы'}<=set(w.sheetnames)
main=w['Все загруженные файлы'];assert main['M5'].value=='Номер тома' and main['M7'].number_format=='@'
assert "!N7:N11" in w['СВОД']['D7'].value,w['СВОД']['D7'].value
work=w['Итоги работы'];assert work['C5'].value=='коробов' and work['D5'].value=='томов'
assert work['C6'].value=='=$H$13' and work['D6'].value=='=$C$13'
assert isinstance(work['B7'].value,datetime) and isinstance(work['B8'].value,datetime)
assert work['C7'].value=='=$H$14' and work['D7'].value=='=$C$14'
assert w['Стат. по типам']['D7'].value=='=C7/100',w['Стат. по типам']['D7'].value
stats=w['Стат. по зам.']
assert [stats[a].font.color.rgb for a in ['C6','E6','I6','M6']]==['FF996600','FF00B050','FFFF0000','FF0070C0']
for tab,a in [('Итоги работы','I13'),('Стат. по зам.','G10')]:
 cell=w[tab][a];assert all(getattr(cell.border,e).style in (None,'none') for e in ['left','right','top','bottom'])
assert w['СВОД']['E7'].border.left.style=='dashed' and w['СВОД']['D7'].border.right.style=='dashed'
for tab in ['СВОД','Стат. по типам','Стат. по зам.','Итоги работы']:assert w[tab]['A4'].border.top.style=='thick'
with zipfile.ZipFile(target) as z:
 for a in z.namelist():
  if a.startswith('xl/worksheets/sheet'):
   d=E.fromstring(z.read(a));assert not any(f.get('t')=='shared' for f in d.iter(n+'f'))
first=work['C6'].value,work['D6'].value
run();w=load_workbook(target);assert (w['Итоги работы']['C6'].value,w['Итоги работы']['D6'].value)==first
assert sum(w['Все загруженные файлы'].cell(5,c).value=='Номер тома' for c in range(1,23))==1
assert sum(bool(v.tabSelected) for sh in w for v in sh.views.sheetView)==1
print('Split report tests passed: four sheets, shared formulas, automatic text column, all colors, outer borders, dashed separators, empty areas and repeat merge')
print(target)
