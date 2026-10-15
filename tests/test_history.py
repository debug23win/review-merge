"""Past days of the reports keep their numbers; volumes by «Номер тома»; recalculation by a separate command. Artificial data only."""
import json, shutil, subprocess, uuid
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXE = ROOT / 'Свод_проверки.exe'
MAIN = 'Все загруженные файлы'
B = ROOT / 'tests/.runs' / uuid.uuid4().hex
B.mkdir(parents=True)


def book(path, documents):
    w = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx'); m = w[MAIN]; m.delete_rows(7, m.max_row)
    for r, doc in enumerate(documents, 7):
        values = {1: r - 6, 4: 'pdf', 6: '01.10.2026', 7: '10:00', 8: f'{r:08X}', 11: 'ПД'}; values.update(doc)
        for c, v in values.items(): m.cell(r, c, v)
    w.save(path); return path


def build(date, *sources):
    """Sources may include switches such as --recalculate-history."""
    subprocess.run([str(EXE), '--batch', str(output), date, MAIN, str(output), *map(str, sources)], check=True, timeout=120)
    return json.loads(Path(str(output) + '.run.json').read_text(encoding='utf-8-sig'))


def checked_volumes(path):
    """«Проверено томов» of every date row of «СВОД» (cached) and whether the cell is a formula."""
    values, formulas = load_workbook(path, data_only=True)['СВОД'], load_workbook(path)['СВОД']
    return {values.cell(r, 2).value.day: (values.cell(r, 10).value, str(formulas.cell(r, 10).value).startswith('='))
            for r in range(8, values.max_row + 1) if isinstance(values.cell(r, 2).value, datetime)}


def reviewer_file(name, edit):
    """A reviewer's copy of the consolidated book with some rows changed; the consolidated book keeps its cached values."""
    path = B / name; shutil.copy(output, path)
    w = load_workbook(path); m = w[MAIN]
    rows = {m.cell(r, 3).value: r for r in range(7, m.max_row + 1) if m.cell(r, 3).value}
    edit(m, rows); w.save(path); return path


documents = [{3: 'Том 1', 9: 'Анна', 10: datetime(2026, 10, 1), 12: 1},
             {3: 'Том 2', 9: 'Борис', 10: datetime(2026, 10, 2), 12: 1},
             {3: 'Том 3', 12: 2},
             {3: 'Раздел 4 книга А', 12: 2}, {3: 'Раздел 4 книга Б', 12: 2},
             {3: 'Том 5', 12: 3}, {3: 'Том 5-УЛ', 12: 3}]
output = book(B / 'svod.xlsx', documents)
build('2026-10-02')
assert checked_volumes(output) == {1: (1, True), 2: (2, True)}, checked_volumes(output)

# A late record for 01.10 arrives on 03.10: the past days keep the numbers already reported, the report date counts it.
late = reviewer_file('late.xlsx', lambda m, rows: (m.cell(rows['Том 3'], 9, 'Анна'), m.cell(rows['Том 3'], 10, datetime(2026, 10, 1))))
build('2026-10-03', late)
assert checked_volumes(output) == {1: (1, False), 2: (2, False), 3: (3, True)}, checked_volumes(output)
# The «Справка» stays live: it already counts the late volume on its day.
# A separate command recalculates every day by the current data and rules.
build('2026-10-03', '--recalculate-history')
assert checked_volumes(output) == {1: (2, True), 2: (3, True), 3: (3, True)}, checked_volumes(output)

# «Номер тома»: files of different names with one number are one volume; a file without the number joins the numbered
# volume of its name (an ИУЛ next to its volume); the log lists the numbers that join different names.
def numbers(m, rows):
    for name, number in (('Раздел 4 книга А', '4'), ('Раздел 4 книга Б', '4'), ('Том 5', '5'), ('Том 1', 'б/н'), ('Том 2', 'б/н')):
        m.cell(rows[name], 13, number)
    for name in ('Раздел 4 книга А', 'Раздел 4 книга Б', 'Том 5', 'Том 5-УЛ'):
        m.cell(rows[name], 9, 'Борис'); m.cell(rows[name], 10, datetime(2026, 10, 4))
report = build('2026-10-04', reviewer_file('numbers.xlsx', numbers))
assert any(note.startswith('«4»:') and 'Раздел 4 книга А' in note and 'Раздел 4 книга Б' in note for note in report['VolumeMerges']), report['VolumeMerges']
assert not any(note.startswith(('«5»', '«б/н»')) for note in report['VolumeMerges']), 'one name is no merge; a bare «б/н» is no number'
# Two volumes more on 04.10: «4» (two files) and «5» (with its ИУЛ); the past days keep their numbers.
assert checked_volumes(output) == {1: (2, False), 2: (3, False), 3: (3, False), 4: (5, True)}, checked_volumes(output)
reference = load_workbook(output, data_only=True)['Справка по томам']
heads = {reference.cell(6, c).value: c for c in range(1, reference.max_column + 1) if reference.cell(6, c).value}
volumes = [reference.cell(r, heads['Том (фрагменты и ИУЛ объединены)']).value for r in range(7, 20)]
assert '№ 4' in volumes and '№ 5' in volumes and 'РАЗДЕЛ 4 КНИГА А' not in volumes and 'ТОМ 5' not in volumes, volumes
print('History tests passed: frozen past days, live report date, recalculation, volumes by number with their IUL and merge notes')
