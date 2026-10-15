"""«Замечания устранены» on the main sheet and the volumes with remarks under «СВОД»; artificial data only."""
import json, shutil, subprocess, uuid
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
EXE = ROOT / 'Свод_проверки.exe'
MAIN = 'Все загруженные файлы'
B = ROOT / 'tests/.runs' / uuid.uuid4().hex
B.mkdir(parents=True)


def book(path, documents):
    w = load_workbook(ROOT / 'tests/fixtures/Анна.xlsx'); m = w[MAIN]; m.delete_rows(7, m.max_row)
    for r, doc in enumerate(documents, 7):
        values = {1: r - 6, 4: 'pdf', 6: '01.10.2026', 7: '10:00', 8: f'{r:08X}', 11: 'ПД'}; values.update(doc)
        for c, v in values.items(): m.cell(r, c, v)
    w.save(path); return path


def build(output, *sources):
    """Sources may include switches such as --remarks-with-borders."""
    subprocess.run([str(EXE), '--batch', str(output), '2026-10-04', MAIN, str(output), *map(str, sources)], check=True, timeout=120)
    return json.loads(Path(str(output) + '.run.json').read_text(encoding='utf-8-sig'))


def rows(path):
    """Main sheet rows by document name."""
    m = load_workbook(path)[MAIN]
    return {m.cell(r, 3).value: r for r in range(7, m.max_row + 1) if m.cell(r, 3).value}


def remarks(path, framed=False):
    """The three values right above the «СВОД» table, two lines each: cached texts, formulas, merges, colors and frames."""
    cached, formulas = load_workbook(path, data_only=True)['СВОД'], load_workbook(path)['СВОД']
    row = next(r for r in range(1, formulas.max_row + 1) if str(cached.cell(r, 4).value or '').startswith('Всего томов\nс замечаниями'))
    assert all(formulas.cell(row, c).alignment.wrap_text for c in (4, 8, 12)) and '&CHAR(10)&' in formulas.cell(row, 4).value, 'two lines in a cell'
    assert row == 4 and formulas.cell(5, 1).value == 'Этап' and formulas.cell(5, 4).value == 'Всего', 'the boxes are right above the table'
    edges = [getattr(formulas.cell(row, c).border, e).style for c in range(4, 16) for e in ('left', 'right', 'top', 'bottom')]
    if framed: assert all(formulas.cell(row, c).border.top.style == 'thin' and formulas.cell(row, c).border.bottom.style == 'thin' for c in range(4, 16)), edges
    else: assert all(e in (None, 'none') for e in edges), ('plain text without frames', edges)
    merges = set(map(str, formulas.merged_cells.ranges))
    for first in (4, 8, 12):
        assert '%s%d:%s%d' % (chr(64 + first), row, chr(64 + first + 3), row) in merges, (first, row)
    colors = [formulas.cell(row, c).font.color.rgb for c in (4, 8, 12)]
    assert colors == ['FF996600', 'FF00B050', 'FFFF0000'] and all(formulas.cell(row, c).font.b for c in (4, 8, 12)), colors
    assert all(formulas.cell(row, c).value.startswith('="') and 'COUNTIFS(' in formulas.cell(row, c).value for c in (4, 8, 12))
    area = formulas.print_area if isinstance(formulas.print_area, str) else ','.join(formulas.print_area)
    assert '$A$1:$S$%d' % formulas.max_row in area.split(',')[0], (area, formulas.max_row)
    return [cached.cell(row, c).value for c in (4, 8, 12)]


documents = [{3: 'Том 1', 9: 'Анна', 10: datetime(2026, 10, 2), 12: 1, 13: 1},
             {3: 'Том 2_фрагмент1', 9: 'Анна', 10: datetime(2026, 10, 2), 12: 1, 19: 'Ошибка описи'},
             {3: 'Том 2_фрагмент2', 9: 'Анна', 10: datetime(2026, 10, 2), 12: 1, 20: 'Примечание'},
             {3: 'Том 3', 9: 'Борис', 10: datetime(2026, 10, 3), 12: 2, 14: 1},
             {3: 'Том 4', 9: 'Борис', 10: datetime(2026, 10, 3), 12: 2},
             {3: 'Том 5', 9: 'Борис', 10: datetime(2026, 10, 5), 12: 3, 13: 1},
             {3: 'Том 6', 12: 3, 19: 'Не проверен'}]
output = book(B / 'svod.xlsx', documents)
build(output)

# The column is added after «Примечание» with the header, its number and an empty cell for every document.
m = load_workbook(output)[MAIN]
assert m['U5'].value == 'Примечание' and m['V5'].value == 'Замечания устранены' and m['V6'].value == 22, [m.cell(5, c).value for c in range(19, 24)]
assert all(m.cell(r, 22).value is None for r in range(7, 14))
assert m.auto_filter.ref is None or m.auto_filter.ref.split(':')[1].startswith('V')
# Checked up to 04.10 with remarks: Том 1, Том 2 (both fragments), Том 3. Том 5 is checked later, Том 6 is not checked.
assert remarks(output) == ['Всего томов\nс замечаниями: 3', 'Устранено замечаний\nпо томам: 0', 'Устраняется томов\nс замечаниями: 3']

# Marks typed in the consolidated book: 1 and a date resolve, «нет» and 0 do not; a volume is resolved when every part with remarks is.
w = load_workbook(output); m = w[MAIN]; r = rows(output)
m.cell(r['Том 1'], 22, 1); m.cell(r['Том 2_фрагмент1'], 22, datetime(2026, 10, 4)); m.cell(r['Том 2_фрагмент2'], 22, 'нет'); m.cell(r['Том 3'], 22, 0)
w.save(output)
build(output)
assert remarks(output) == ['Всего томов\nс замечаниями: 3', 'Устранено замечаний\nпо томам: 1', 'Устраняется томов\nс замечаниями: 2']
m = load_workbook(output)[MAIN]
assert m.cell(r['Том 2_фрагмент1'], 22).value == datetime(2026, 10, 4) and m.cell(r['Том 2_фрагмент1'], 22).is_date, 'a date stays a date, its cell keeps its format'
assert m.cell(r['Том 1'], 22).value == 1 and m.cell(r['Том 2_фрагмент2'], 22).value == 'нет'

# A reviewer's file with the column updates the consolidated book: a filled value of a later file wins, an empty one changes nothing.
reviewer = B / 'Борис.xlsx'; shutil.copy(output, reviewer)
w = load_workbook(reviewer); m = w[MAIN]
m.cell(r['Том 2_фрагмент2'], 22, 1); m.cell(r['Том 3'], 22, 'да'); m.cell(r['Том 1'], 22, None); m.cell(r['Том 5'], 22, datetime(2026, 10, 6))
w.save(reviewer)
build(output, reviewer)
m = load_workbook(output)[MAIN]
assert [m.cell(r[name], 22).value for name in ('Том 1', 'Том 2_фрагмент2', 'Том 3')] == [1, 1, 'да']
assert m.cell(r['Том 5'], 22).value == datetime(2026, 10, 6) and m.cell(r['Том 5'], 22).number_format.lower() == 'dd.mm.yyyy', 'a new date gets the date format'
assert remarks(output) == ['Всего томов\nс замечаниями: 3', 'Устранено замечаний\nпо томам: 3', 'Устраняется томов\nс замечаниями: 0']

# The reference sheet keeps the rule as formulas: an open remark per file and a resolved flag per volume.
reference = load_workbook(output)['Справка по томам']
heads = {reference.cell(6, c).value: c for c in range(1, reference.max_column + 1) if reference.cell(6, c).value}
assert 'Замечание не устранено' in heads and 'Замечания тома устранены' in heads, sorted(h for h in heads if isinstance(h, str))[:5]
open_cell = reference.cell(7, heads['Замечание не устранено']).value
assert open_cell.startswith('=IF(') and "'Все загруженные файлы'!V" in open_cell and '"НЕТ"' in open_cell, open_cell
assert reference.cell(7, heads['Замечания тома устранены']).value.startswith('=IF(AND(')

# Plain text by default; the option frames the same values.
build(output, '--remarks-with-borders')
assert remarks(output, framed=True) == ['Всего томов\nс замечаниями: 3', 'Устранено замечаний\nпо томам: 3', 'Устраняется томов\nс замечаниями: 0']

# «Итоги работы»: non-zero boxes and volumes of the first table are bold, zeros are not; rules keep it so in Excel.
work, values = load_workbook(output)['Итоги работы'], load_workbook(output, data_only=True)['Итоги работы']
last = 6  # the first table ends at its last date row; the alternative table below has dates too
while isinstance(values.cell(last + 1, 2).value, datetime): last += 1
checked = 0
for r in range(6, last + 1):
    for c in range(3, work.max_column + 1):
        v = values.cell(r, c).value
        if isinstance(v, (int, float)):
            assert bool(work.cell(r, c).font.b) == (v != 0), (r, c, v, work.cell(r, c).font.b); checked += 1
assert checked >= 20 and any(values.cell(r, c).value == 0 for r in range(7, last + 1) for c in range(3, 8))
rules = [(rule.operator, rule.formula) for cf in work.conditional_formatting for rule in cf.rules]
assert ('notEqual', ['0']) in rules and ('equal', ['0']) in rules, rules
total = next(c for c in range(3, work.max_column + 1) if work.cell(4, c).value == 'ВСЕГО')
assert any(str(cf.sqref) == 'C6:%s%d' % (get_column_letter(total + 1), last) for cf in work.conditional_formatting), [str(cf.sqref) for cf in work.conditional_formatting]

# The work totals title is not merged; the PDF shows the boxes on the first page.
work = load_workbook(output)['Итоги работы']
assert not any(str(x).startswith('A1:') for x in work.merged_cells.ranges) and work['A1'].alignment.horizontal == 'left'
report = json.loads(Path(str(output) + '.run.json').read_text(encoding='utf-8-sig'))
assert report['PdfError'] is None and Path(report['Pdf']).exists(), report
print('Remarks summary tests passed: resolved column, merge from files, per-volume rule, two-line values above the summary table with or without frames, formulas, bold non-zero work totals and unmerged title')
