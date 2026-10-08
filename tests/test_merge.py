"""Regression checks against two artificial XLSX review copies."""
import json
import hashlib
import shutil
from openpyxl import load_workbook
from copy import copy
from datetime import datetime
import subprocess
import uuid
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def cells(archive, member):
    doc = ET.fromstring(archive.read(member))
    result = {}
    for cell in doc.findall('.//s:sheetData/s:row/s:c', NS):
        value = ''.join(t.text or '' for t in cell.findall('.//s:t', NS)) if cell.get('t') == 'inlineStr' else cell.findtext('s:v', '', NS)
        result[cell.get('r')] = (value, cell.findtext('s:f', None, NS))
    return result

def run(output, sources, current_date=False, *flags):
    options=(['--current-date-for-new'] if current_date else [])+list(flags)
    subprocess.run([str(ROOT/'Свод_проверки.exe'), '--batch', str(output), *options, '2026-10-04', 'Все загруженные файлы', *map(str, sources)], check=True, timeout=90)
    return json.loads(Path(str(output)+'.run.json').read_text(encoding='utf-8-sig'))

run_root = ROOT/'tests/.runs'
run_root.mkdir(exist_ok=True)
def test_merge():
    temp = run_root/uuid.uuid4().hex
    temp.mkdir()
    result = temp/'result.xlsx'
    report = run(result, [ROOT/'tests/fixtures/Анна.xlsx', ROOT/'tests/fixtures/Борис.xlsx'])
    assert (report['Documents'], report['Reviewed'], report['WithIssues'], report['Boxes']) == (7, 6, 4, 3), report
    assert report['Conflicts'] == 11 and report['Problems'] == 3, report
    with zipfile.ZipFile(result) as z:
        wb = ET.fromstring(z.read('xl/workbook.xml'))
        names = [s.get('name') for s in wb.findall('s:sheets/s:sheet', NS)]
        assert set(names) == {'Все загруженные файлы', 'Свод', 'Справка по томам', 'Служебный лист'}, names
        main = cells(z, 'xl/worksheets/sheet1.xml')
        notes = [v for addr, (v, f) in main.items() if addr.startswith(('S','T'))]
        assert any('=1+1' in v for v in notes)
        assert all(f is None for addr, (v, f) in main.items() if addr.startswith(('S','T')))
        assert not any(addr.startswith(('U','V','W','X','Y','Z','AA')) for addr,(v,f) in main.items() if v or f)
        assert any('COUNTIFS' in f for member in z.namelist() if member.startswith('xl/worksheets/') and member.endswith('.xml') for v,f in cells(z,member).values() if f)
        assert not any('[Анна,' in v or '[Борис,' in v for v in notes)
        for addr in ['M7','N7','Q9']:
            assert main[addr][0] == '1', (addr, main[addr])
    repeated = temp/'repeated.xlsx'
    rerun = run(repeated, [result, ROOT/'tests/fixtures/Анна.xlsx', ROOT/'tests/fixtures/Борис.xlsx'])
    assert rerun['Documents'] == 7 and rerun['WithIssues'] == 4
    with zipfile.ZipFile(result) as a, zipfile.ZipFile(repeated) as b:
        first = cells(a, 'xl/worksheets/sheet1.xml')
        second = cells(b, 'xl/worksheets/sheet1.xml')
        for addr in ['S7','T7','AA7']:
            assert first.get(addr) == second.get(addr), (addr, first.get(addr), second.get(addr))
def test_existing_target():
    temp=run_root/uuid.uuid4().hex;temp.mkdir()
    target=temp/'existing.xlsx';shutil.copyfile(ROOT/'tests/fixtures/Анна.xlsx',target)
    wb=load_workbook(target);ws=wb['Все загруженные файлы']
    ws['A7']='=ROW()-5';ws['E7']='=100+23';ws['AB7']='Мои данные'
    ws['I100']='Строка без документа';ws['T100']='Сохранить это замечание'
    wb.create_sheet('Ручные итоги')['A1']='=2+3'
    wb.create_sheet('История проверок')['A1']='Личная история — сохранить'
    wb.save(target)
    before=target.read_bytes()
    source=ROOT/'tests/fixtures/Борис.xlsx';source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
    report=run(target,[source,target,source])
    assert report['Documents']==7 and report['WithIssues']==4,report
    backups=list(temp.glob('existing.xlsx.backup-*'));assert len(backups)==1
    assert backups[0].read_bytes()==before
    assert hashlib.sha256(source.read_bytes()).hexdigest()==source_hash
    wb=load_workbook(target);ws=wb['Все загруженные файлы']
    assert ws['A7'].value=='=ROW()-5' and ws['E7'].value=='=100+23'
    assert ws['AB7'].value=='Мои данные' and ws['I100'].value=='Строка без документа'
    assert ws['T100'].value=='Сохранить это замечание'
    assert ws['M7'].value==1 and ws['N7'].value==1
    assert wb['Ручные итоги']['A1'].value=='=2+3'
    assert wb['История проверок']['A1'].value=='Личная история — сохранить'
    assert 'Свод' in wb.sheetnames and 'Конфликты свода' not in wb.sheetnames
    assert sum(bool(ws.cell(r,3).value) for r in range(101,ws.max_row+1))==2
    notes=ws['S7'].value,ws['T7'].value
    run(target,[source]);again=load_workbook(target)['Все загруженные файлы']
    assert notes==(again['S7'].value,again['T7'].value)
    assert again['A7'].value=='=ROW()-5'

def test_user_columns():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();target=temp/'occupied.xlsx'
    shutil.copyfile(ROOT/'tests/fixtures/Анна.xlsx',target)
    wb=load_workbook(target);wb['Все загруженные файлы']['U7']='Сохранить';wb.save(target)
    before=target.read_bytes()
    args=[str(ROOT/'Свод_проверки.exe'),'--batch',str(target),'2026-10-04','Все загруженные файлы',str(ROOT/'tests/fixtures/Борис.xlsx')]
    assert subprocess.run(args,timeout=90).returncode==0
    assert load_workbook(target)['Все загруженные файлы']['U7'].value=='Сохранить'
    assert list(temp.glob('*.backup-*'))[0].read_bytes()==before

def test_template_and_fragments():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();source=temp/'template.xlsx'
    shutil.copyfile(ROOT/'tests/fixtures/Анна.xlsx',source)
    wb=load_workbook(source);main=wb['Все загруженные файлы']
    main['C7']='Любой шифр A/7 Фрагмент 1';main['C8']='Любой шифр A/7 Фрагмент 2'
    for r in [7,8]:
        main.cell(r,9,'Проверяющий');main.cell(r,10,datetime(2026,10,2));main.cell(r,13,1)
    main['T7']='Первое исходное замечание';main['T8']='Второе исходное замечание'
    main['S7']='1. Ошибка в наименовании\n2. Ошибка в номере изменения';main['S8']=main['S7'].value
    main['U7']='Пользовательские данные U';main['AA7']='Пользовательские данные AA'
    for sheet in list(wb):
        if sheet.title.lower()=='свод':wb.remove(sheet)
    ref=wb.create_sheet('СВОД');ref.merge_cells('A1:D1');ref['A1']='СВОДНАЯ ТАБЛИЦА';ref['A2']='Мой объект';ref['A7']='Получено коробов';ref['A8']='Кол-во томов по акту'
    physical=wb.create_sheet('Справка по томам');physical['A1']='Коробов получено';physical['B1']=50;physical['B2']=800;physical['B8']=datetime(2026,10,2);physical['G42']=5;physical['G43']=120
    for c,date in [(2,datetime(2026,10,2)),(3,datetime(2026,10,4)),(4,datetime(2026,10,5))]:
        ref.cell(5,c,date);ref.cell(5,c).number_format='dd.mm.yyyy'
        ref.cell(7,c,"='Справка по томам'!B1");ref.cell(8,c,"='Справка по томам'!B2")
    ref['B9']="='Справка по томам'!G42";ref['B10']="='Справка по томам'!G43"
    ref['A17']='Сохранить исходное примечание';ref['A65']='Пользовательский текст ниже статистики'
    ref['A7'].font=copy(main['C7'].font);style=copy(ref['A7']._style)
    wb.save(source);output=temp/'result.xlsx';run(output,[source])
    w=load_workbook(output);s=w['Свод'];m=w['Все загруженные файлы']
    assert s['A1'].value=='СВОДНАЯ ТАБЛИЦА' and s['A2'].value=='Мой объект'
    assert s['A7']._style==style and s['A17'].value=='Сохранить исходное примечание'
    assert s['A65'].value=='Пользовательский текст ниже статистики'
    assert 'Справка по томам' in s['B10'].value
    assert m['U7'].value=='Пользовательские данные U' and m['AA7'].value=='Пользовательские данные AA'
    assert m['T7'].value=='Первое исходное замечание' and m['T8'].value=='Второе исходное замечание'
    p=w['Справка по томам'];marker=next(c for c in p[1] if c.value=='ReviewMerge.FirstDates.v3');vh=marker.column+12
    assert not any(c.value=='ReviewMerge.FirstDates.v3' for c in s[1])
    assert p['B1'].value==50 and p['B2'].value==800
    matching=[r for r in range(7,p.max_row+1) if p.cell(r,vh).value=='ЛЮБОЙ ШИФР A/7']
    assert len(matching)==1,matching
    vc=p.cell(matching[0],vh+3);assert vc.value.startswith('=MAX(')
    cached=load_workbook(output,data_only=True)['Справка по томам'];assert cached.cell(matching[0],vh+9).value==1,'A filled S field is one remark per volume'
    assert cached.cell(matching[0],vh+10).value==1
    before_notes=m['T7'].value,m['T8'].value;run(output,[source]);s2=load_workbook(output)['Свод'];m2=load_workbook(output)['Все загруженные файлы']
    assert before_notes==(m2['T7'].value,m2['T8'].value)
    assert 'Справка по томам' in s2['B10'].value
    assert sum(bool(v.tabSelected) for sh in load_workbook(output) for v in sh.views.sheetView)==1

def test_legacy_annotations():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();source=temp/'legacy.xlsx';shutil.copyfile(ROOT/'tests/fixtures/Анна.xlsx',source)
    wb=load_workbook(source);ws=wb['Все загруженные файлы']
    headers=['Ключ короба (формула)','Первая строка короба','Первая проверка короба','Документ проверен','Дата для статистики','Есть замечания','Первая дата из источников']
    for c,text in enumerate(headers,21):ws.cell(5,c,text)
    ws['AA7']=datetime(2026,10,1);ws['T7']='[Анна, 01.10.2026] Первый текст\n\n[Борис, 02.10.2026] Второй текст'
    wb.save(source);output=temp/'clean.xlsx';run(output,[source])
    ws=load_workbook(output)['Все загруженные файлы']
    assert ws['T7'].value=='Первый текст\n\nВторой текст'
    assert all(ws.cell(5,c).value is None and ws.cell(7,c).value is None for c in range(21,28))
    run(output,[source]);assert load_workbook(output)['Все загруженные файлы']['T7'].value=='Первый текст\n\nВторой текст'

def test_complete_boxes_and_reviewers():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();source=temp/'boxes.xlsx'
    wb=load_workbook(ROOT/'tests/fixtures/Анна.xlsx');main=wb['Все загруженные файлы']
    main.delete_rows(7,main.max_row)
    documents=[('Том А_фрагмент1','Анна',2,1),('Том А Фрагмент 2','Анна',2,1),('Том Б',None,None,1),('Том В_изм.1','Борис',3,2),('Том Г','Борис',4,2),('Том Д',None,None,None),('Том А-УЛ','Анна',2,1),('Том В-УЛ','Борис',3,2)]
    for r,(file,who,day,box) in enumerate(documents,7):
        for c,value in {1:r-6,3:file,4:'pdf',8:f'{r:08X}',9:who,10:datetime(2026,10,day) if day else None,11:'ПД',12:box}.items():main.cell(r,c,value)
    main['J7']='02.10.2026';main['J7'].number_format='@'
    from openpyxl.utils.datetime import to_excel
    main['J8']=to_excel(datetime(2026,10,2));main['J8'].number_format='General'
    main.sheet_view.tabSelected=True
    other=wb['Служебный лист'];other.sheet_view.tabSelected=True
    wb.save(source);output=temp/'result.xlsx';run(output,[source])
    formulas=load_workbook(output);cached=load_workbook(output,data_only=True)
    ref=formulas['Справка по томам'];summary=cached['Свод'];marker=next(c for c in ref[1] if c.value=='ReviewMerge.FirstDates.v3')
    start=cached['Справка по томам'].cell(1,marker.column+3).value
    assert cached['Справка по томам'].cell(start,3).value==2
    assert cached['Справка по томам'].cell(start,4).value==1
    assert cached['Справка по томам'].cell(start,6).value==datetime(2026,10,2)
    assert cached['Справка по томам'].cell(start+1,6).value==datetime(2026,10,3)
    names={summary.cell(r,1).value:r for r in range(51,54)}
    assert summary.cell(names['Анна'],2).value==1 and summary.cell(names['Борис'],2).value==2
    assert summary.cell(names['Анна'],5).value==3
    assert sum(bool(sh.sheet_view.tabSelected) for sh in formulas)==1
    assert formulas.active.title=='Свод'
    assert str(formulas['Свод'].print_area).endswith('$'+str(formulas['Свод'].max_row))
    assert cached['Все загруженные файлы']['J7'].value==datetime(2026,10,2)
    assert cached['Все загруженные файлы']['J8'].value==datetime(2026,10,2)
    assert formulas['Все загруженные файлы']['J7'].number_format=='dd.mm.yyyy'
    assert formulas['Все загруженные файлы']['J8'].number_format=='dd.mm.yyyy'
    ref.cell(start,5,2);ref.cell(start+1,5,2);formulas.save(output)
    run(output,[source]);after=load_workbook(output,data_only=True)['Справка по томам']
    assert after.cell(start,5).value==2 and after.cell(start+1,5).value==2
    assert after.cell(start,6).value==datetime(2026,10,2) and after.cell(start+1,6).value==datetime(2026,10,3)
    assert sum(c.value=='ReviewMerge.FirstDates.v3' for c in load_workbook(output)['Справка по томам'][1])==1

def test_current_date_for_new_reviews():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();source=temp/'incoming.xlsx'
    w=load_workbook(ROOT/'tests/fixtures/Анна.xlsx');m=w['Все загруженные файлы'];m.delete_rows(7,m.max_row)
    documents=[('Том А','Анна',datetime(2026,10,2)),('Том Б','Борис',datetime(2026,10,3)),('Том В','Борис',None),('Том Г',None,None),('Том Д','Анна','неверная дата'),('Том А Фрагмент 2','Анна',datetime(2026,10,2))]
    for r,(file,who,date) in enumerate(documents,7):
        for c,v in {1:r-6,3:file,4:'pdf',8:f'{r:08X}',9:who,10:date,11:'ПД',12:1,20:'Исходный текст'}.items():m.cell(r,c,v)
    w.save(source)
    target=temp/'current.xlsx';shutil.copyfile(source,target)
    w=load_workbook(target);m=w['Все загруженные файлы']
    for r in range(8,13):m.cell(r,9).value=None;m.cell(r,10).value=None
    w.save(target)
    today=datetime.now().replace(hour=0,minute=0,second=0,microsecond=0)
    report=run(target,[source],True);assert report['DatesAssignedToday']==4,report
    w=load_workbook(target,data_only=True);m=w['Все загруженные файлы']
    assert m['J7'].value==datetime(2026,10,2)
    for r in [8,9,11,12]:assert m.cell(r,10).value==today,(r,m.cell(r,10).value)
    assert m['J10'].value is None and m['I10'].value is None
    assert all(m.cell(r,20).value=='Исходный текст' for r in range(7,13))
    assert all(m.cell(r,10).number_format=='dd.mm.yyyy' for r in range(7,13))
    ref=w['Справка по томам'];h=next(c.column for c in ref[1] if c.value=='ReviewMerge.FirstDates.v3')
    assert ref.cell(8,h+6).value==today,'Source historical date leaked into first-review statistics'
    before=[m.cell(r,10).value for r in range(7,13)];repeat=run(target,[source],True)
    after=load_workbook(target,data_only=True)['Все загруженные файлы'];assert repeat['DatesAssignedToday']==0
    assert before==[after.cell(r,10).value for r in range(7,13)],'Repeated import changed saved dates'
    fresh=temp/'fresh.xlsx';run(fresh,[source],True);n=load_workbook(fresh,data_only=True)['Все загруженные файлы']
    assert all(n.cell(r,10).value==today for r in [7,8,9,11,12])
    original=temp/'original-dates.xlsx';run(original,[source]);n=load_workbook(original,data_only=True)['Все загруженные файлы']
    assert n['J7'].value==datetime(2026,10,2) and n['J8'].value==datetime(2026,10,3)
    assert n['J9'].value is None and n['J11'].value=='неверная дата'

MAIN='Все загруженные файлы'
def book(path, documents, people=None):
    """Artificial register: one row per document, columns given by number (C=3, I=9, J=10, L=12, ...)."""
    w=load_workbook(ROOT/'tests/fixtures/Анна.xlsx');m=w[MAIN];m.delete_rows(7,m.max_row)
    for r,doc in enumerate(documents,7):
        values={1:r-6,4:'pdf',6:'01.10.2026',7:'10:00',8:f'{r:08X}',11:'ПД'};values.update(doc)
        for c,v in values.items():m.cell(r,c,v)
    if people is not None:
        ref=w.create_sheet('Справка по томам');ref['A1']='Коробов получено';ref['B1']=50;ref['A2']='Томов по проекту акта';ref['B2']=100;ref['A8']='ФИО проверяющего'
        for n,name in enumerate(people):ref.cell(12+2*n,1,name)
    w.save(path);return path

def summary_column(sheet, day, header_row=35):
    return next(c.column for c in sheet[header_row] if c.value==day)

def test_box_lists_and_dates_with_year():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();source=temp/'boxes.xlsx'
    book(source,[{3:'Том 1',9:'Анна',10:'02.10.2026 г.',12:'45 и 46'},{3:'Том 2',9:'Анна',10:datetime(2026,10,2),12:'47,48'},
                 {3:'Том 3',9:'Анна',10:datetime(2026,10,2),12:'50-52'},{3:'Том 4',9:'Анна',10:datetime(2026,10,2),12:3},
                 {3:'Том 5',12:'60 и 61'},{3:'Том 6',9:'Анна',10:datetime(2026,10,2),12:'к. №7'},{3:'Том 7',9:'Анна',10:datetime(2026,10,2),12:'без номера'}])
    output=temp/'result.xlsx';report=run(output,[source])
    assert report['Boxes']==9 and report['Problems']==1,report
    cached=load_workbook(output,data_only=True);main=cached[MAIN];ref=cached['Справка по томам']
    assert main['J7'].value==datetime(2026,10,2) and load_workbook(output)[MAIN]['J7'].number_format=='dd.mm.yyyy'
    marker=next(c for c in ref[1] if c.value=='ReviewMerge.FirstDates.v3');start=ref.cell(1,marker.column+3).value;end=ref.cell(1,marker.column+4).value
    boxes=[ref.cell(r,2).value for r in range(start,end+1)]
    assert boxes==[3,7,45,46,47,48,50,51,52,60,61],boxes
    dated={ref.cell(r,2).value:ref.cell(r,6).value for r in range(start,end+1)}
    assert dated[45]==dated[46]==datetime(2026,10,2) and dated[60] in (None,'')
    slots=[c.value for row in ref.iter_rows(min_row=12,max_row=12) for c in row if isinstance(c.value,str)]
    assert '45пд, 46пд' in slots,slots

def test_surname_matching():
    temp=run_root/uuid.uuid4().hex;temp.mkdir()
    people=['Яковлева Александра Алексеевна','Шекунов Виталий Викторович']
    source=book(temp/'source.xlsx',[{3:'Том 1',9:'АА Яковлева',10:datetime(2026,10,2),12:1},{3:'Том 2',9:'АА Яковлева',10:datetime(2026,10,2),12:1},
                                     {3:'Том 3',9:'шекунов',10:datetime(2026,10,2),12:2},{3:'Том 4',9:'Шекунов',10:datetime(2026,10,2),12:2},{3:'Том 5',9:'Белякова',10:datetime(2026,10,2),12:3}])
    for flags,expected in [(['--match-surnames'],None),([],'АА Яковлева')]:
        target=book(temp/f'target{len(flags)}.xlsx',[{3:f'Том {n}'} for n in range(1,6)],people)
        report=run(target,[source],False,*flags)
        cached=load_workbook(target,data_only=True);ref=cached['Справка по томам'];summary=cached['Свод']
        rows={ref.cell(r,1).value for r in range(12,60,2)}-{None}
        assert people[0] in rows and people[1] in rows and 'шекунов' not in rows and 'Шекунов' not in rows,rows
        assert (expected in rows)==(expected is not None),rows
        labels={summary.cell(r,1).value:summary.cell(r,5).value for r in range(51,54)}
        assert labels=={'Белякова':1,'Шекунов':2,'Яковлева':2},labels
        if flags:assert any('АА Яковлева' in m and people[0] in m for m in report['PeopleMerges']),report

def test_incomplete_boxes_and_standalone_iul():
    temp=run_root/uuid.uuid4().hex;temp.mkdir()
    documents=[{3:'Том А',9:'Анна',10:datetime(2026,10,2),12:5},{3:'Том А-УЛ'},{3:'Том Б Фрагмент 1',9:'Анна',10:datetime(2026,10,2),12:5},{3:'Том Б Фрагмент 2'},
               {3:'Том В',9:'Анна',10:datetime(2026,10,2),12:6},{3:'Том Г',12:5},{3:'960-СМ2-УЛ-1',9:'Анна',10:datetime(2026,10,2)},{3:'960-СМ2-УЛ-2'},
               {3:'Том Д',9:'Анна',10:datetime(2026,10,2)},{3:'Том Д-УЛ'}]
    source=book(temp/'source.xlsx',documents)
    plain=run(temp/'plain.xlsx',[source]);filled=run(temp/'filled.xlsx',[source],False,'--fill-incomplete')
    assert len(plain['IncompleteBoxes'])==2 and plain['FilledRows']==0,plain
    assert plain['IncompleteBoxes'][0].startswith('Короб 5пд: отмечено 2 из 5 строк') and plain['IncompleteBoxes'][1].startswith('Том без номера короба Том Д'),plain['IncompleteBoxes']
    assert filled['IncompleteBoxes']==[] and filled['FilledRows']==4 and filled['Reviewed']==9,filled
    for name,volumes,files in [('plain.xlsx',1,5),('filled.xlsx',5,9)]:
        summary=load_workbook(temp/name,data_only=True)['Свод'];col=summary_column(summary,datetime(2026,10,4))
        assert summary.cell(38,col).value==volumes and summary.cell(37,col).value==files,(name,summary.cell(38,col).value,summary.cell(37,col).value)
    main=load_workbook(temp/'filled.xlsx',data_only=True)[MAIN]
    assert main['I8'].value=='Анна' and main['J8'].value==datetime(2026,10,2) and main['L8'].value==5
    assert main['I12'].value=='Анна' and main['L12'].value==5,'Remaining rows of the box are filled too'
    assert main['I14'].value is None,'A standalone ИУЛ is not a part of another volume'

def test_filled_rows_keep_reviewer_and_box():
    temp=run_root/uuid.uuid4().hex;temp.mkdir()
    target=book(temp/'svod.xlsx',[{3:'Том 1',9:'Яковлева А.А.',10:datetime(2026,10,1),12:8,19:'1. Первое'}])
    source=book(temp/'reviewer.xlsx',[{3:'Том 1',9:'Иванов',10:datetime(2026,10,3),12:9,13:1,19:'1. Первое\n2. Второе'}])
    run(target,[source]);main=load_workbook(target,data_only=True)[MAIN]
    assert (main['I7'].value,main['J7'].value,main['L7'].value)==('Яковлева А.А.',datetime(2026,10,1),8)
    assert main['M7'].value==1 and main['S7'].value=='1. Первое\n2. Второе'

def test_kind_table_and_notes():
    temp=run_root/uuid.uuid4().hex;temp.mkdir();target,source=temp/'svod.xlsx',temp/'reviewer.xlsx'
    for path,s,t in [(target,'1. Ошибка А','Первый абзац'),(source,'1. Ошибка А\n2. Ошибка Б','Первый абзац\n\nВторой абзац')]:
        w=load_workbook(ROOT/'tests/fixtures/Анна.xlsx');w[MAIN]['S7']=s;w[MAIN]['T7']=t;w.save(path)
    first=run(target,[source]);second=run(target,[source])
    main=load_workbook(target)[MAIN]
    assert main['S7'].value=='1. Ошибка А\n2. Ошибка Б' and main['T7'].value=='Первый абзац\n\nВторой абзац',(main['S7'].value,main['T7'].value)
    assert first['Conflicts']==2 and second['Conflicts']==0,(first['Conflicts'],second['Conflicts'])
    summary=load_workbook(target,data_only=True)['Свод']
    assert summary['A48'].value=='Не проверено' and all(summary.cell(r,5).value is None for r in range(44,49))
    assert summary['B48'].value==first['Documents']-first['Reviewed']+1,'«вчера» is not a date, so that reviewed file is not checked'

def test_long_calendar_filter_duplicates():
    temp=run_root/uuid.uuid4().hex;temp.mkdir()
    documents=[{3:f'Документ-{i}',9:'Анна',10:datetime(2025,10,6) if i==1 else datetime(2026,10,2),12:i//4+1} for i in range(40)]
    output=temp/'calendar.xlsx';run(output,[book(temp/'source.xlsx',documents)])
    w=load_workbook(output);summary=w['Свод'];ref=w['Справка по томам']
    columns=[c.column for c in summary[5] if c.value is not None and c.column>=2]
    assert columns==list(range(2,max(columns)+1)),'Date columns of the summary must not be cut by calculation columns'
    assert next(c.column for c in ref[1] if c.value=='ReviewMerge.FirstDates.v3')>max(columns)
    assert all("'Справка по томам'!" in summary.cell(37,c).value for c in columns if summary.cell(37,c).value)
    target,source=temp/'svod.xlsx',temp/'reviewer.xlsx'
    w=load_workbook(ROOT/'tests/fixtures/Анна.xlsx');m=w[MAIN];m.auto_filter.ref='A5:V11';m['V5']='Мой столбец'
    for c in range(1,21):m.cell(20,c,m.cell(7,c).value)
    w.save(target)
    w=load_workbook(ROOT/'tests/fixtures/Анна.xlsx');w[MAIN]['P7']=1;w.save(source)
    run(target,[source]);m=load_workbook(target)[MAIN]
    assert m.auto_filter.ref.startswith('A5:V'),m.auto_filter.ref
    assert m['P7'].value==1 and m['P20'].value==1
    with zipfile.ZipFile(target) as z:assert 'forceFullCalc' not in z.read('xl/workbook.xml').decode()

def test_marks_parts_and_wide_days():
    temp=run_root/uuid.uuid4().hex;temp.mkdir()
    documents=[{3:f'Том {n}',9:'Анна',10:datetime(2026,10,2),12:n} for n in range(1,8)]
    documents+=[{3:'Том Ч_Часть1',9:'Анна',10:datetime(2026,10,3),12:9,14:'да'},{3:'Том Ч (Часть 2)',9:'Анна',10:datetime(2026,10,3),12:9,15:'+'}]
    source=book(temp/'source.xlsx',documents)
    target=book(temp/'svod.xlsx',[{3:d[3]} for d in documents],['Анна'])
    w=load_workbook(target);ref=w['Справка по томам'];ref['B8']=datetime(2026,10,3);ref['G9']='ИТОГО';ref['H12']='Моё примечание';w.save(target)
    plain=run(temp/'plain.xlsx',[source]);assert plain['Problems']==2 and plain['ConvertedMarks']==[],plain
    report=run(target,[source],False,'--convert-marks');assert report['Problems']==0 and len(report['ConvertedMarks'])==2,report
    main=load_workbook(target)[MAIN];assert main['N14'].value==1 and main['O15'].value==1
    cached=load_workbook(target,data_only=True);ref=cached['Справка по томам']
    days=[(c.column,c.value) for c in ref[8] if isinstance(c.value,datetime)]
    assert days[0]==(2,datetime(2026,10,2)) and days[1][1]==datetime(2026,10,3),days
    slots=[ref.cell(12,c).value for c in range(2,days[1][0])]
    assert [v for v in slots if isinstance(v,str) and v.endswith('пд')]==[f'{n}пд' for n in range(1,8)],slots
    assert ref.cell(12,days[1][0]+6).value=='Моё примечание','The note moves together with its day'
    summary=cached['Свод'];col=summary_column(summary,datetime(2026,10,4))
    assert summary.cell(38,col).value==8,'«Часть 1» and «Часть 2» are one volume'

test_merge()
test_existing_target()
test_user_columns()
test_template_and_fragments()
test_legacy_annotations()
test_complete_boxes_and_reviewers()
test_current_date_for_new_reviews()
test_box_lists_and_dates_with_year()
test_surname_matching()
test_incomplete_boxes_and_standalone_iul()
test_filled_rows_keep_reviewer_and_box()
test_marks_parts_and_wide_days()
test_kind_table_and_notes()
test_long_calendar_filter_duplicates()
print('Merge tests passed: union, flags, comments, conflicts, invalid data, formulas, repeated merge, box lists, surnames, incomplete volumes and ИУЛ')

